from pathlib import Path
from zipfile import ZipFile
import re, json
from xml.etree import ElementTree as ET
from pypdf import PdfReader
import openpyxl

BASE=Path(r"D:\GED_PROFISSIONAL\tmp\enterprise_r12")
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}

with ZipFile(BASE/'deck_ref.pptx') as z:
    slides=[]
    for n in sorted([x for x in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',x)], key=lambda x:int(re.search(r'\d+',x).group())):
        root=ET.fromstring(z.read(n)); texts=[t.text or '' for t in root.findall('.//a:t',NS)]
        slides.append({'slide':n,'texts':texts})
    (BASE/'deck_text.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')

reader=PdfReader(str(BASE/'relatorio_ref.pdf'))
(BASE/'pdf_text.txt').write_text('\n\n'.join(f'--- PAGE {i+1} ---\n{p.extract_text() or ""}' for i,p in enumerate(reader.pages)),encoding='utf-8')

wb=openpyxl.load_workbook(BASE/'hist_ref.xlsx',data_only=False,read_only=False)
out=[]
for ws in wb.worksheets:
    rows=[]
    for row in ws.iter_rows():
        vals=[c.value for c in row]
        if any(v not in (None,'') for v in vals): rows.append({'row':row[0].row,'values':vals[:20]})
    out.append({'sheet':ws.title,'max_row':ws.max_row,'max_col':ws.max_column,'rows':rows[:120]})
(BASE/'workbook_text.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print('slides',len(slides),'pdf pages',len(reader.pages),'sheets',wb.sheetnames)
