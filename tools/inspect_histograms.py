import json
from pathlib import Path
import openpyxl

FILES = [
    Path(r"D:\Doc Control\Histogramas\HISTOGRAMA_Doc_Control_18-08-26_R10.xlsx"),
    Path(r"D:\Doc Control\Histogramas\HISTOGRAMA_Arquivo Técnico + Doc Control _20-08-26_R10.xlsx"),
]

for path in FILES:
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=False)
    print("\nFILE", path.name)
    for ws_f in wb_f.worksheets:
        if ws_f.title.startswith("HIST MOI"):
            ws_v = wb_v[ws_f.title]
            print(f"\nSHEET {ws_f.title}")
            for r in [10, 11, 12, 13, 14, 15, 23, 24, 25, 26, 34, 35, 36, 37, 46, 47, 48, 49, 50, 51]:
                populated = []
                for c in range(1, 76):
                    cv = ws_v.cell(r, c).value
                    if cv is not None:
                        populated.append([ws_v.cell(r, c).coordinate, cv])
                if populated:
                    print(json.dumps({"row": r, "cells": populated}, ensure_ascii=False, default=str))
            continue
        if ws_f.title not in ("00 - RESUMO PARA DECIS�O", "AN�LISE REDU��O", "CUSTOS EXECUTIVOS"):
            continue
        ws_v = wb_v[ws_f.title]
        print(f"\nSHEET {ws_f.title} DIM {ws_f.max_row}x{ws_f.max_column}")
        for row in ws_f.iter_rows():
            populated = []
            for c in row:
                if c.value is not None:
                    populated.append({"cell": c.coordinate, "formula": c.value, "value": ws_v[c.coordinate].value})
            if populated:
                print(json.dumps(populated, ensure_ascii=False, default=str))
