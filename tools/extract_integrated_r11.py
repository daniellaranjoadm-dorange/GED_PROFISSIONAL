import json
from pathlib import Path
import openpyxl

integrated = Path(r"D:\Doc Control\Histogramas\HISTOGRAMA_Arquivo Técnico + Doc Control _20-08-26_R10.xlsx")
archive = next(p for p in Path(r"D:\Doc Control\Histogramas").rglob("HISTOGRAMA_Arquivo *_20-08-26_R10.xlsx") if "+" not in p.name)

def vals(ws, row, start=8, count=68):
    return [ws.cell(row, c).value or 0 for c in range(start, start + count)]

wi = openpyxl.load_workbook(integrated, data_only=True)
wa = openpyxl.load_workbook(archive, data_only=True)
sa = wi["HIST MOI_Acordado com Borghesan"]
sd = wi["HIST MOI_Doc Control"]
st = wi["HIST MOI Arquivo Técnico"]

data = {
    "periods": [f"M{i}" for i in range(1, 69)],
    "dates": [str(sa.cell(1, c).value or "") for c in range(8, 76)],
    "rates": {
        "Gestão Doc Control": 11000,
        "Supervisor Arquivo Técnico": 11850,
        "Técnico Documentação": 8122.053,
        "Assistente Documentação": 5884.768,
        "Auxiliar Documentação": 4710,
        "Analista Qualidade": 5548.168,
    },
    "agreed": {
        "total": vals(sa, 47),
        "projects": {"Handy Size": vals(sa, 10), "Gaseiros / GLP": vals(sa, 24), "Mid Range / MR1": vals(sa, 37)},
        "hm": {"Handy Size": sa["F10"].value, "Gaseiros / GLP": sa["F24"].value, "Mid Range / MR1": sa["F37"].value},
    },
    "doc_control": {
        "total": vals(sd, 43),
        "projects": {"Handy Size": vals(sd, 10), "Gaseiros / GLP": vals(sd, 23), "Mid Range / MR1": vals(sd, 34)},
        "roles": {
            "Gestão Doc Control": {"Handy Size": vals(sd, 11), "Gaseiros / GLP": vals(sd, 24), "Mid Range / MR1": vals(sd, 35)},
            "Supervisor Arquivo Técnico": {"Handy Size": vals(sd, 12), "Gaseiros / GLP": [0]*68, "Mid Range / MR1": [0]*68},
            "Técnico Documentação": {"Handy Size": vals(sd, 13), "Gaseiros / GLP": vals(sd, 25), "Mid Range / MR1": vals(sd, 36)},
            "Assistente Documentação": {"Handy Size": vals(sd, 14), "Gaseiros / GLP": vals(sd, 26), "Mid Range / MR1": vals(sd, 37)},
        },
    },
    "archive": {
        "total": vals(st, 39),
        "projects": {"Handy Size": vals(st, 10), "Gaseiros / GLP": vals(st, 21), "Mid Range / MR1": vals(st, 31)},
        "roles": {
            "Assistente Documentação": {"Handy Size": vals(st, 11), "Gaseiros / GLP": vals(st, 22), "Mid Range / MR1": vals(st, 32)},
            "Auxiliar Documentação": {"Handy Size": vals(st, 12), "Gaseiros / GLP": vals(st, 23), "Mid Range / MR1": vals(st, 33)},
            "Analista Qualidade": {"Handy Size": vals(st, 13), "Gaseiros / GLP": vals(st, 24), "Mid Range / MR1": vals(st, 34)},
        },
    },
}
out = Path(r"D:\GED_PROFISSIONAL\tmp\integrated_r11\source_data.json")
out.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
print(out, out.stat().st_size)
