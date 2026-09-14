from pathlib import Path
import openpyxl, json

p=Path(r"D:\Doc Control\Histogramas\PACOTE_DASHBOARD_DIRETORIA_R12_ENTERPRISE\HISTOGRAMA_INTEGRADO_DOC_CONTROL_ARQUIVO_TECNICO_R13_ENTERPRISE.xlsx")
wf=openpyxl.load_workbook(p,data_only=False,read_only=False)
wv=openpyxl.load_workbook(p,data_only=True,read_only=False)
print('SHEETS',wf.sheetnames)
keys=['RESUMO','CUST','ARQUIVO','REDU','SAV','PREMISS','CHECK','HIST']
for ws in wf.worksheets:
    print('\n###',ws.title,ws.max_row,ws.max_column)
    shown=0
    for r in range(1,ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1,min(ws.max_column,18)+1)]
        txt=' | '.join('' if v is None else str(v) for v in vals)
        if any(k in txt.upper() for k in keys) or (ws.title.upper().find('ARQUIVO')>=0 and any(v not in (None,'') for v in vals)):
            dvals=[wv[ws.title].cell(r,c).value for c in range(1,min(ws.max_column,18)+1)]
            print(r,txt[:600], ' || DATA:', dvals[:12])
            shown+=1
            if shown>=80: break
