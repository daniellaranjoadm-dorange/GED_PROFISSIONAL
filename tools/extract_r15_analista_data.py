from pathlib import Path
import openpyxl,json,glob
src=Path(glob.glob(r"D:\GED_PROFISSIONAL\tmp\enterprise_r15_analista\**\*.xlsx",recursive=True)[0])
base=Path(r"D:\GED_PROFISSIONAL\tmp\integrated_r11\source_data.json")
out=Path(r"D:\GED_PROFISSIONAL\tmp\enterprise_r15_analista\source_data.json")
data=json.loads(base.read_text(encoding='utf-8'));w=openpyxl.load_workbook(src,data_only=True);ws=w['HIST_ARQUIVO_TECNICO'];projects=['Handy Size','Gaseiros / GLP','Mid Range / MR1']
roles={}
for role,rows in {'Assistente Documentação':[5,6,7],'Auxiliar Documentação':[8,9,10]}.items():roles[role]={p:[ws.cell(r,c).value or 0 for c in range(7,75)] for p,r in zip(projects,rows)}
shared=[ws.cell(11,c).value or 0 for c in range(7,75)]
roles['Analista Qualidade']={'Handy Size':shared,'Gaseiros / GLP':[0]*68,'Mid Range / MR1':[0]*68}
data['archive']['roles']=roles;data['archive']['projects']={p:[sum(roles[r][p][i] for r in roles) for i in range(68)] for p in projects};data['archive']['total']=[sum(data['archive']['projects'][p][i] for p in projects) for i in range(68)];data['archive']['shared_analyst_label']='Handy Size / GLP / MR1';data['archive']['shared_analyst_hm']=sum(shared)
out.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');print({'archive_hm':sum(data['archive']['total']),'combined_hm':sum(data['archive']['total'])+sum(data['doc_control']['total']),'archive_peak':max(data['archive']['total']),'combined_peak':max(a+b for a,b in zip(data['archive']['total'],data['doc_control']['total'])),'analyst_shared_hm':sum(shared)})
