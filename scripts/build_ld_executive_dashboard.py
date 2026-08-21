"""Gera o Dashboard Executivo da LD Projeto Básico em um único HTML."""

from __future__ import annotations

import base64
import json
import re
import sys
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SOURCE = Path(
    r"\\virm-rgr022\FILESERVER\Projetos\05_HANDYMAX\09. Doc Control\3 - LD"
    r"\I-LD-4880.00-9311-000-CZ1-001_RD.xlsm"
)
DEFAULT_OUTPUT = Path(
    r"C:\Users\daniel.laranjo\.codex\visualizations\2026\08\03"
    r"\019fc763-53de-7740-999c-b52eea40f680\outputs\ld_executive_dashboard"
    r"\dashboard_executivo_ld_projeto_basico.html"
)
DEFAULT_LOGO = Path(r"D:\GED_PROFISSIONAL\Logo icone de Pasta.png")
LD_MAX_COLUMN = 61  # A:BI


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value in (None, ""):
        return None
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def integer(value):
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def date_or_text(value):
    """Preserva datas como ISO e ocorrências operacionais como REPROVADO."""
    if isinstance(value, time):
        return None
    converted = iso(value)
    if converted:
        return converted
    text = str(value or "").strip()
    return text or None


def revision_rank(value):
    text = str(value if value is not None else "0").strip().upper() or "0"
    if text.isdigit():
        return (0, int(text))
    return (1, sum((ord(char) - 64) * (27 ** pos) for pos, char in enumerate(reversed(text)) if "A" <= char <= "Z"))


def _normalized_text(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^A-Z0-9]+", " ", "".join(char for char in text if not unicodedata.combining(char)).upper()).strip()


def _worksheet_by_tokens(workbook, *tokens):
    wanted = tuple(_normalized_text(token) for token in tokens)
    for sheet in workbook.worksheets:
        title = _normalized_text(sheet.title)
        if all(token in title for token in wanted):
            return sheet
    return None


def _http_link(cell):
    target = ""
    if cell.hyperlink is not None:
        target = str(cell.hyperlink.target or "").strip()
    if not target:
        target = str(cell.value or "").strip()
    return target if re.match(r"^https?://", target, flags=re.IGNORECASE) else ""


def load_dashboard_links(source: Path):
    """Le links DOX/PCF das abas auxiliares sem alterar a planilha fonte."""
    workbook = load_workbook(source, read_only=False, data_only=False, keep_vba=False, keep_links=True)
    dox_links = {}
    pcf_links = {}

    fap = _worksheet_by_tokens(workbook, "FAP", "PRODU")
    if fap is not None:
        for row in range(2, fap.max_row + 1):
            document = _normalized_text(fap.cell(row, 2).value).replace(" ", "")
            revision = _normalized_text(fap.cell(row, 3).value).replace(" ", "") or "0"
            link = _http_link(fap.cell(row, 1))
            if document and link:
                dox_links[(document, revision)] = link

    pcf_sheet = _worksheet_by_tokens(workbook, "LISTA", "DOCUMENT", "PCF")
    if pcf_sheet is not None:
        # Exportacao DOX: Nome em B, Rotulo/Revisao em D e URL em L.
        for row in range(2, pcf_sheet.max_row + 1):
            name = str(pcf_sheet.cell(row, 2).value or "").strip()
            link = _http_link(pcf_sheet.cell(row, 12))
            identifier = _normalized_text(name).replace(" ", "")
            if not identifier.startswith("PCF") or not link:
                continue
            pcf_links[identifier] = link

    workbook.close()
    return dox_links, pcf_links


def clean_pcf_status(value):
    text = str(value or "").strip().upper()
    return {"NOT RELESED": "NOT RELEASED", "NOT RELEASE": "NOT RELEASED"}.get(text, text or "SEM PCF")


def export_cell(value):
    """Converte uma celula da LD para JSON sem perder numeros ou datas."""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return value


def load_ld_headers(source: Path):
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    sheet = workbook["LD PROJETO BASICO"]
    headers = [str(cell.value or "").strip() for cell in sheet[1][:LD_MAX_COLUMN]]
    workbook.close()
    return headers


def load_records(source: Path):
    dox_links, pcf_links = load_dashboard_links(source)
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    sheet = workbook["LD PROJETO BASICO"]
    grouped = {}
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=LD_MAX_COLUMN, values_only=True), 2
    ):
        document = str(values[2] or "").strip()
        if not document or document.upper() in {"NOT APPLICABLE", "N/A", "#N/A"}:
            continue
        revision = values[3] if values[3] not in (None, "") else 0
        record = {
            "documento": document,
            "tipo": str(values[0] or "-").strip(),
            "revisao": str(revision),
            "titulo": str(values[4] or "-").replace("\n", " ").strip(),
            "disciplina": str(values[6] or values[5] or "SEM DISCIPLINA").strip(),
            "especialidade": str(values[14] or "-").strip(),
            "status": str(values[12] or "SEM STATUS").strip(),
            "statusEmissao": str(values[13] or "SEM STATUS").strip(),
            "grd": str(values[22] or "").strip(),
            "dataEmissao": iso(values[23]),
            "pcf": str(values[24] or "").strip(),
            "dataPcf": iso(values[25]),
            "statusPcf": clean_pcf_status(values[26]),
            "pcfRespondida": str(values[27] or "").strip(),
            "dataResposta": iso(values[28]),
            "grdResposta": str(values[29] or "").strip(),
            "responsavel": str(values[21] or "SEM RESPONSÁVEL").strip(),
            "dataProgramada": None,
            "pcpMarenova": "",
            "guiaEmissao": str(values[34] or "").strip(),
            "statusDox": str(values[46] or "SEM REGISTRO").strip(),
            "documentoKm": str(values[7] or "").strip(),
            "kmTitle": str(values[8] or "").replace("\n", " ").strip(),
            "transmittalKm": str(values[9] or "").strip(),
            "dataKm": iso(values[10]),
            "documentoKmEmitidoDocTp": str(values[11] or "").strip(),
            "medicaoEmissao": date_or_text(values[47]),
            "medicaoAprovacao": date_or_text(values[48]),
            "casco": str(values[54] or "").strip(),
            "linkDox": dox_links.get(
                (_normalized_text(document).replace(" ", ""), _normalized_text(revision).replace(" ", "") or "0"),
                "",
            ),
            # Vínculo documental estrito: a URL precisa pertencer exatamente à
            # PCF exibida (incluindo revisão), nunca a uma revisão semelhante.
            "linkPcf": pcf_links.get(
                _normalized_text(values[24]).replace(" ", ""),
                "",
            ),
            "open": integer(values[31]),
            "comentarios": integer(values[30]),
            "underReview": integer(values[32]),
            "statusFinalPcf": clean_pcf_status(values[33]),
            "cronogramaInicio": iso(values[55]),
            "cronogramaTermino": iso(values[56]),
            "cronogramaInicioOriginal": date_or_text(values[55]),
            "cronogramaTerminoOriginal": date_or_text(values[56]),
            "bmEmissao": str(values[57] or "SEM INFORMAÇÃO").strip(),
            "bmDataEmissao": date_or_text(values[58]),
            "bmAprovacao": str(values[59] or "SEM INFORMAÇÃO").strip(),
            "bmDataAprovacao": date_or_text(values[60]),
            "linhaFonte": row_number,
            "ldExport": [export_cell(value) for value in values],
        }
        key = (document, str(revision).strip())
        grouped.setdefault(key, []).append(record)

    consolidated = []
    status_priority = {"Emitido": 5, "Recebido e não Emitido": 4, "Reprovado": 3, "Aprovado com comentários": 3, "Aguardando PCF": 2, "Recebido": 2, "Não Recebido": 1}
    for records in grouped.values():
        records.sort(key=lambda row: (status_priority.get(row["status"], 0), bool(row["pcf"]), bool(row["dataKm"])), reverse=True)
        merged = dict(records[0])
        merged["documentoKm"] = " | ".join(dict.fromkeys(row["documentoKm"] for row in records if row["documentoKm"]))
        merged["kmTitle"] = " | ".join(dict.fromkeys(row["kmTitle"] for row in records if row["kmTitle"]))
        merged["transmittalKm"] = " | ".join(dict.fromkeys(row["transmittalKm"] for row in records if row["transmittalKm"]))
        merged["documentoKmEmitidoDocTp"] = " | ".join(dict.fromkeys(row["documentoKmEmitidoDocTp"] for row in records if row["documentoKmEmitidoDocTp"]))
        for field in (
            "dataKm", "medicaoEmissao", "medicaoAprovacao", "cronogramaInicio",
            "cronogramaTermino", "cronogramaInicioOriginal", "cronogramaTerminoOriginal",
            "bmEmissao", "bmDataEmissao", "bmAprovacao", "bmDataAprovacao",
        ):
            merged[field] = next((row[field] for row in records if row.get(field)), None)
        merged["open"] = max(row["open"] for row in records)
        merged["comentarios"] = max(row["comentarios"] for row in records)
        merged["underReview"] = max(row["underReview"] for row in records)
        # Uma linha por documento/revisao, preservando todos os vinculos KM.
        merged["ldExport"][7] = merged["documentoKm"]
        merged["ldExport"][8] = merged["kmTitle"]
        merged["ldExport"][9] = merged["transmittalKm"]
        merged["ldExport"][10] = merged["dataKm"]
        merged["ldExport"][11] = merged["documentoKmEmitidoDocTp"]
        consolidated.append(merged)

    latest = {}
    for record in consolidated:
        document, revision = record["documento"], record["revisao"]
        current = latest.get(document)
        if current is None or revision_rank(revision) >= revision_rank(current[0]):
            latest[document] = (revision, record["linhaFonte"], record)
    workbook.close()
    return [item[2] for item in latest.values()]


HTML = r'''<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Command Center | LD Projeto Básico</title>
<script src="https://cdn.tailwindcss.com"></script><script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.7/dist/chart.umd.min.js"></script><script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
<script>tailwind.config={theme:{extend:{colors:{ink:'#06131e',panel:'#0d2233',line:'#234359',cyan:'#37c8f4',mint:'#1dd6b5',amber:'#ffad32',danger:'#ff5e68'}}}}</script>
<style>:root{color-scheme:dark}*{box-sizing:border-box}body{min-height:100vh;background:radial-gradient(circle at 8% 0,#153b53 0,#071722 34%,#030c14 100%)}.glass{background:linear-gradient(145deg,rgba(13,34,51,.96),rgba(6,19,30,.97));border:1px solid rgba(102,180,220,.19);box-shadow:0 18px 44px #0005}.hero-brand{min-height:110px;min-width:min(760px,100%);display:flex;align-items:center;gap:1.15rem}.brand-logo{flex:0 0 auto;width:78px;height:78px;object-fit:cover;object-position:center;border-radius:18px;border:1px solid rgba(255,158,27,.38);box-shadow:0 10px 28px rgba(0,0,0,.38),0 0 0 5px rgba(255,158,27,.045)}.hero-copy{text-shadow:0 2px 15px #03101bd9}.control,.multi-btn{width:100%;border:1px solid #294d64;background:#071925;color:#eef8fc;border-radius:.75rem;padding:.7rem .8rem;outline:none}.control:focus,.multi.open .multi-btn{border-color:#37c8f4;box-shadow:0 0 0 3px #37c8f422}.multi{position:relative}.multi-btn{display:flex;align-items:center;justify-content:space-between;gap:.5rem;text-align:left}.multi-btn span:first-child{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.multi-panel{display:none;position:absolute;z-index:80;top:calc(100% + 6px);left:0;min-width:100%;width:max-content;max-width:360px;max-height:330px;overflow:auto;border:1px solid #315b73;border-radius:.8rem;background:#071925;box-shadow:0 20px 45px #000a;padding:.45rem}.multi.open .multi-panel{display:block}.multi-option{display:flex;align-items:center;gap:.6rem;min-width:210px;padding:.55rem .65rem;border-radius:.5rem;color:#dcebf2;cursor:pointer}.multi-option:hover{background:#123047}.multi-option:first-child{border-bottom:1px solid #294d64;margin-bottom:.25rem;font-weight:800;color:#37c8f4}.multi-option input{width:16px;height:16px;accent-color:#37c8f4}.kpi{position:relative;overflow:hidden}.kpi:before{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:var(--a)}.chart{height:330px}.scroll::-webkit-scrollbar{height:8px}.scroll::-webkit-scrollbar-thumb{background:#294d64;border-radius:10px}@media(max-width:700px){.hero-brand{min-height:96px;gap:.8rem}.brand-logo{width:62px;height:62px;border-radius:14px}}@media print{body{background:#fff}.glass{background:#fff;color:#111;box-shadow:none}.no-print{display:none!important}}</style></head>
<body class="font-sans text-slate-100 antialiased"><main class="mx-auto max-w-[1700px] px-4 py-5 md:px-8">
<header class="mb-5 flex flex-col gap-4 xl:flex-row xl:items-end xl:justify-between"><div class="hero-brand"><img class="brand-logo" src="__LOGO__" alt=""><div class="hero-copy"><p class="mb-2 text-xs font-black uppercase tracking-[.2em] text-cyan">Naval Engineering Intelligence</p><h1 class="text-3xl font-black md:text-5xl">LD Project Command Center</h1><p class="mt-2 text-sm text-slate-400">Visão executiva consolidada · última revisão de cada documento · LD Projeto Básico</p></div></div><div class="flex flex-wrap gap-2 text-xs"><span class="rounded-full border border-line bg-panel px-3 py-2">Fonte: I-LD-4880.00-9311-000-CZ1-001</span><span id="base" class="rounded-full border border-mint/30 bg-mint/10 px-3 py-2 font-bold text-mint"></span><span id="scope" class="rounded-full border border-cyan/30 bg-cyan/10 px-3 py-2 font-bold text-cyan"></span></div></header>

<section class="glass no-print mb-5 rounded-2xl p-5"><div class="mb-4 flex items-center justify-between"><div><h2 class="font-black">Filtros gerenciais</h2><p class="text-xs text-slate-400">Todos os KPIs, gráficos e a base detalhada são atualizados em conjunto.</p></div><button id="clear" class="rounded-lg border border-line px-3 py-2 text-xs font-bold hover:border-cyan hover:text-cyan">Limpar filtros</button></div><div class="grid gap-3 sm:grid-cols-2 lg:grid-cols-4 xl:grid-cols-7">
<div class="text-xs font-bold text-slate-400">Tipo<div id="tipo" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Disciplina<div id="disciplina" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Status geral<div id="status" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Emissão<div id="emissao" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Prazo do cronograma<div id="prazo" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Medição Emissão<div id="medEmissao" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Medição Aprovação<div id="medAprovacao" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Status PCF<div id="pcfStatus" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Responsável<div id="responsavel" class="multi mt-1"></div></div><div class="text-xs font-bold text-slate-400">Casco<div id="casco" class="multi mt-1"></div></div></div>
<div class="mt-3 grid gap-3 lg:grid-cols-[1fr_auto]"><input id="search" class="control" placeholder="Buscar documento, título, GRD, PCF ou KM..."><button id="export" class="rounded-xl bg-cyan px-5 py-3 text-sm font-black text-ink hover:brightness-110">Exportar recorte Excel</button></div></section>

<section class="mb-5 grid gap-4 sm:grid-cols-2 xl:grid-cols-4"><article class="glass kpi rounded-2xl p-5" style="--a:#37c8f4"><p class="text-xs font-black uppercase tracking-wider text-cyan">Documentos no escopo</p><div id="kTotal" class="mt-3 text-4xl font-black">—</div><p id="kTotalSub" class="mt-2 text-xs text-slate-400"></p></article><article class="glass kpi rounded-2xl p-5" style="--a:#1dd6b5"><p class="text-xs font-black uppercase tracking-wider text-mint">Progresso de emissão</p><div id="kEmitidos" class="mt-3 text-4xl font-black">—</div><p id="kEmitidosSub" class="mt-2 text-xs text-slate-400"></p></article><article class="glass kpi rounded-2xl p-5" style="--a:#ffad32"><p class="text-xs font-black uppercase tracking-wider text-amber">Recebidos não emitidos</p><div id="kBacklog" class="mt-3 text-4xl font-black">—</div><p id="kBacklogSub" class="mt-2 text-xs text-slate-400"></p></article><article class="glass kpi rounded-2xl p-5" style="--a:#ff5e68"><p class="text-xs font-black uppercase tracking-wider text-danger">PCFs críticas</p><div id="kPcf" class="mt-3 text-4xl font-black">—</div><p id="kPcfSub" class="mt-2 text-xs text-slate-400"></p></article></section>
<section id="insight" class="glass mb-5 rounded-2xl border-l-4 border-l-amber p-5"></section>
<section class="grid gap-5 xl:grid-cols-2"><article class="glass rounded-2xl p-5"><p class="text-xs font-black uppercase tracking-wider text-cyan">Carteira</p><h2 class="mb-4 text-lg font-black">Situação documental</h2><div class="chart"><canvas id="statusChart"></canvas></div></article><article class="glass rounded-2xl p-5"><p class="text-xs font-black uppercase tracking-wider text-mint">Produção</p><h2 class="mb-4 text-lg font-black">Emitidos x pendentes por disciplina</h2><div class="chart"><canvas id="disciplineChart"></canvas></div></article><article class="glass rounded-2xl p-5"><p class="text-xs font-black uppercase tracking-wider text-amber">Composição</p><h2 class="mb-4 text-lg font-black">Volume por tipo de documento</h2><div class="chart"><canvas id="typeChart"></canvas></div></article><article class="glass rounded-2xl p-5"><p class="text-xs font-black uppercase tracking-wider text-danger">Qualidade e aprovação</p><h2 class="mb-4 text-lg font-black">Status das PCFs recebidas</h2><div class="chart"><canvas id="pcfChart"></canvas></div></article></section>
<section class="glass mt-5 rounded-2xl p-5"><div class="mb-4 flex flex-col gap-2 md:flex-row md:items-end md:justify-between"><div><p class="text-xs font-black uppercase tracking-wider text-danger">Gestão por exceção</p><h2 class="text-xl font-black">Documentos prioritários</h2><p class="text-xs text-slate-400">Recebidos ainda não emitidos, vencidos no cronograma, PCFs não liberadas e comentários abertos aparecem primeiro.</p></div><span id="tableCount" class="text-xs font-bold text-slate-400"></span></div><div class="scroll overflow-x-auto"><table class="w-full min-w-[1750px] text-left text-sm"><thead class="border-b border-line text-[11px] uppercase tracking-wider text-slate-400"><tr><th class="p-3">Documento / título</th><th class="p-3">Rev.</th><th class="p-3">Tipo</th><th class="p-3">Disciplina</th><th class="p-3">Situação</th><th class="p-3">Emissão</th><th class="p-3">Início cronograma</th><th class="p-3">Término cronograma</th><th class="p-3">Prazo</th><th class="p-3">Medição Emissão</th><th class="p-3">Medição Aprovação</th><th class="p-3">Status PCF</th><th class="p-3 text-right">Open</th><th class="p-3">Responsável</th></tr></thead><tbody id="rows" class="divide-y divide-line/60"></tbody></table></div></section>
<section class="glass mt-5 rounded-2xl p-5"><p class="text-xs font-black uppercase tracking-wider text-amber">Confiabilidade da informação</p><h2 class="mt-1 text-xl font-black">Controles automáticos de qualidade</h2><div id="quality" class="mt-4 grid gap-3 md:grid-cols-2 xl:grid-cols-4"></div></section>
<footer class="py-6 text-center text-xs text-slate-500">D'OR@NGE GED Enterprise · HTML regenerado após a atualização da LD Projeto Básico</footer></main>
<script>
const DATA=__DATA__, META=__META__, $=id=>document.getElementById(id), fmt=new Intl.NumberFormat('pt-BR'), pct=v=>`${(v||0).toLocaleString('pt-BR',{maximumFractionDigits:1})}%`, esc=s=>String(s??'').replace(/[&<>'"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;',"'":'&#39;','"':'&quot;'}[c]));let charts={},filtered=[];Chart.defaults.color='#aac0cc';Chart.defaults.borderColor='#52799133';
const filterMap={tipo:'tipo',disciplina:'disciplina',status:'status',emissao:'statusEmissao',medEmissao:'medicaoEmissao',medAprovacao:'medicaoAprovacao',pcfStatus:'statusPcf',responsavel:'responsavel',casco:'casco'};
const brDate=v=>/^\d{4}-\d{2}-\d{2}$/.test(v||'')?new Date(v+'T12:00:00').toLocaleDateString('pt-BR'):v||'-';
function prazoStatus(r){if(r.statusEmissao==='Emitido')return 'Emitido';if(!r.cronogramaTermino)return 'Cronograma não definido';return r.cronogramaTermino<META.generated.slice(0,10)?'Vencido não emitido':'A vencer não emitido'}
function values(k){return [...new Set(DATA.map(r=>r[k]||'-'))].sort((a,b)=>String(a).localeCompare(String(b),'pt-BR'))}
function checked(id){return new Set([...$(id).querySelectorAll('input[data-value]:checked')].map(x=>x.dataset.value))}
function updateMultiLabel(id){const chosen=checked(id),button=$(id).querySelector('.multi-label'),total=$(id).querySelectorAll('input[data-value]').length;button.textContent=!chosen.size?'Nenhum selecionado':chosen.size===total?'Todos':`${chosen.size} selecionados`}
function initMulti(id,options){const el=$(id);el.innerHTML=`<button type="button" class="multi-btn"><span class="multi-label">Todos</span><span>⌄</span></button><div class="multi-panel"><label class="multi-option"><input type="checkbox" data-all checked> Selecionar tudo</label>${options.map(v=>`<label class="multi-option"><input type="checkbox" data-value="${esc(v)}" checked> <span>${esc(brDate(v))}</span></label>`).join('')}</div>`;const all=el.querySelector('[data-all]'),items=[...el.querySelectorAll('[data-value]')];el.querySelector('.multi-btn').addEventListener('click',e=>{e.preventDefault();document.querySelectorAll('.multi.open').forEach(x=>x!==el&&x.classList.remove('open'));el.classList.toggle('open')});all.addEventListener('change',()=>{items.forEach(x=>x.checked=all.checked);updateMultiLabel(id);update()});items.forEach(x=>x.addEventListener('change',()=>{all.checked=items.every(i=>i.checked);all.indeterminate=!all.checked&&items.some(i=>i.checked);updateMultiLabel(id);update()}));updateMultiLabel(id)}
function init(){Object.entries(filterMap).forEach(([id,k])=>initMulti(id,values(k)));initMulti('prazo',['Vencido não emitido','A vencer não emitido','Cronograma não definido','Emitido']);$('base').textContent=`Atualizado: ${new Date(META.generated).toLocaleString('pt-BR')}`}
function count(rows,k){return rows.reduce((a,r)=>(a[r[k]||'-']=(a[r[k]||'-']||0)+1,a),{})}function selected(){const q=$('search').value.trim().toLocaleLowerCase('pt-BR'),sets=Object.fromEntries(Object.keys(filterMap).map(id=>[id,checked(id)])),deadlines=checked('prazo');return DATA.filter(r=>Object.entries(filterMap).every(([id,k])=>sets[id].has(String(r[k]||'-')))&&deadlines.has(prazoStatus(r))&&(!q||[r.documento,r.titulo,r.grd,r.pcf,r.documentoKm].some(v=>String(v||'').toLocaleLowerCase('pt-BR').includes(q))))}
function draw(id,type,data,options={}){charts[id]?.destroy();charts[id]=new Chart($(id),{type,data,options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{labels:{usePointStyle:true,boxWidth:8}},tooltip:{backgroundColor:'#06131e',borderColor:'#294d64',borderWidth:1,padding:12}},...options}})}
function kpis(r){const n=r.length,emit=r.filter(x=>x.statusEmissao==='Emitido').length,back=r.filter(x=>x.status==='Recebido e não Emitido'||(x.statusEmissao==='Não Emitido'&&x.dataKm)).length,pcfs=r.filter(x=>x.pcf).length,critical=r.filter(x=>x.pcf&&x.statusPcf==='NOT RELEASED').length,open=r.reduce((s,x)=>s+x.open,0);$('kTotal').textContent=fmt.format(n);$('kTotalSub').textContent=`${fmt.format(r.filter(x=>x.status==='Não Recebido').length)} ainda não recebidos`;$('kEmitidos').textContent=pct(n?emit/n*100:0);$('kEmitidosSub').textContent=`${fmt.format(emit)} emitidos · ${fmt.format(n-emit)} pendentes`;$('kBacklog').textContent=fmt.format(back);$('kBacklogSub').textContent='material disponível aguardando emissão';$('kPcf').textContent=fmt.format(critical);$('kPcfSub').textContent=`${fmt.format(open)} comentários OPEN em ${fmt.format(pcfs)} PCFs`;$('scope').textContent=`${fmt.format(n)} de ${fmt.format(DATA.length)} documentos`;$('insight').innerHTML=`<p class="text-xs font-black uppercase tracking-wider text-amber">Leitura executiva</p><p class="mt-2 text-lg font-bold">${n?`A emissão atingiu ${pct(emit/n*100)}. Há ${fmt.format(back)} documentos recebidos aguardando emissão e ${fmt.format(critical)} PCFs classificadas como NOT RELEASED.`:'Nenhum documento corresponde aos filtros.'}</p><p class="mt-2 text-sm text-slate-400">Prioridade: eliminar o backlog já recebido e tratar PCFs não liberadas com comentários abertos.</p>`}
function graphs(r){const st=count(r,'status'),sl=Object.keys(st),palette=['#526c7c','#37c8f4','#ffad32','#ff5e68','#9a83ff','#1dd6b5','#2e8dd0'];draw('statusChart','doughnut',{labels:sl,datasets:[{data:Object.values(st),backgroundColor:palette,borderColor:'#0d2233',borderWidth:4}]},{cutout:'64%',plugins:{legend:{position:'bottom'}}});const ds={};r.forEach(x=>{const k=x.disciplina||'-';ds[k]??={e:0,p:0};x.statusEmissao==='Emitido'?ds[k].e++:ds[k].p++});const d=Object.entries(ds).sort((a,b)=>(b[1].e+b[1].p)-(a[1].e+a[1].p)).slice(0,10).reverse();draw('disciplineChart','bar',{labels:d.map(x=>x[0]),datasets:[{label:'Emitidos',data:d.map(x=>x[1].e),backgroundColor:'#1dd6b5'},{label:'Pendentes',data:d.map(x=>x[1].p),backgroundColor:'#526c7c'}]},{indexAxis:'y',scales:{x:{stacked:true,beginAtZero:true},y:{stacked:true,grid:{display:false},ticks:{font:{size:10}}}}});const ty=Object.entries(count(r,'tipo')).sort((a,b)=>b[1]-a[1]);draw('typeChart','bar',{labels:ty.map(x=>x[0]),datasets:[{label:'Documentos',data:ty.map(x=>x[1]),backgroundColor:'#37c8f4',borderRadius:7}]},{plugins:{legend:{display:false}},scales:{y:{beginAtZero:true,ticks:{precision:0}},x:{grid:{display:false}}}});const p=r.filter(x=>x.pcf),ps=count(p,'statusPcf'),pl=Object.keys(ps);draw('pcfChart','doughnut',{labels:pl,datasets:[{data:Object.values(ps),backgroundColor:pl.map(x=>x==='NOT RELEASED'?'#ff5e68':x==='RELEASED'?'#1dd6b5':x==='RELEASED WITH COMMENTS'?'#ffad32':'#526c7c'),borderColor:'#0d2233',borderWidth:4}]},{cutout:'62%',plugins:{legend:{position:'bottom'}}})}
function priority(x){return (prazoStatus(x)==='Vencido não emitido'?20000:0)+(x.status==='Recebido e não Emitido'?10000:0)+(x.statusPcf==='NOT RELEASED'?5000:0)+x.open*10+(x.underReview||0)}function table(r){const top=[...r].sort((a,b)=>priority(b)-priority(a)).slice(0,30);$('tableCount').textContent=`Exibindo ${top.length} de ${r.length}`;$('rows').innerHTML=top.map(x=>`<tr class="hover:bg-cyan/5"><td class="p-3"><div class="font-bold">${esc(x.documento)}</div><div class="max-w-[430px] truncate text-xs text-slate-500" title="${esc(x.titulo)}">${esc(x.titulo)}</div></td><td class="p-3 font-bold">${esc(x.revisao)}</td><td class="p-3 text-cyan">${esc(x.tipo)}</td><td class="p-3 text-xs">${esc(x.disciplina)}</td><td class="p-3">${esc(x.status)}</td><td class="p-3">${esc(x.statusEmissao)}</td><td class="p-3">${esc(brDate(x.cronogramaInicio))}</td><td class="p-3">${esc(brDate(x.cronogramaTermino))}</td><td class="p-3 ${prazoStatus(x)==='Vencido não emitido'?'font-bold text-danger':''}">${esc(prazoStatus(x))}</td><td class="p-3">${esc(brDate(x.medicaoEmissao))}</td><td class="p-3">${esc(brDate(x.medicaoAprovacao))}</td><td class="p-3 ${x.statusPcf==='NOT RELEASED'?'font-bold text-danger':''}">${esc(x.statusPcf)}</td><td class="p-3 text-right font-black text-amber">${fmt.format(x.open)}</td><td class="p-3 text-xs">${esc(x.responsavel)}</td></tr>`).join('')||'<tr><td colspan="14" class="p-10 text-center text-slate-500">Nenhum registro.</td></tr>'}
function quality(r){const tests=[['Disciplina não mapeada',r.filter(x=>x.disciplina==='#N/A').length,'Corrigir cadastro/mapeamento'],['Emitido sem GRD',r.filter(x=>x.statusEmissao==='Emitido'&&!x.grd).length,'Rastreabilidade incompleta'],['PCF sem data de recebimento',r.filter(x=>x.pcf&&!x.dataPcf).length,'SLA não calculável'],['Resposta sem data',r.filter(x=>x.pcfRespondida&&!x.dataResposta).length,'Resposta não auditável']];$('quality').innerHTML=tests.map(([t,n,s])=>`<article class="rounded-xl border ${n?'border-danger/30 bg-danger/5':'border-mint/30 bg-mint/5'} p-4"><div class="text-2xl font-black ${n?'text-danger':'text-mint'}">${fmt.format(n)}</div><div class="mt-1 font-bold">${t}</div><div class="mt-1 text-xs text-slate-400">${s}</div></article>`).join('')}
function update(){filtered=selected();kpis(filtered);graphs(filtered);table(filtered);quality(filtered)}$('search').addEventListener('input',update);$('clear').addEventListener('click',()=>{[...Object.keys(filterMap),'prazo'].forEach(id=>{const all=$(id).querySelector('[data-all]'),items=[...$(id).querySelectorAll('[data-value]')];all.checked=true;all.indeterminate=false;items.forEach(x=>x.checked=true);updateMultiLabel(id)});$('search').value='';update()});document.addEventListener('click',e=>{if(!e.target.closest('.multi'))document.querySelectorAll('.multi.open').forEach(x=>x.classList.remove('open'))});$('export').addEventListener('click',()=>{const rows=filtered.map(x=>({'Documento':x.documento,'Revisão':x.revisao,'Título':x.titulo,'Tipo':x.tipo,'Disciplina':x.disciplina,'Status':x.status,'Status Emissão':x.statusEmissao,'Início Cronograma':x.cronogramaInicio,'Término Cronograma':x.cronogramaTermino,'Prazo do Cronograma':prazoStatus(x),'Medição Emissão':x.medicaoEmissao,'Medição Aprovação':x.medicaoAprovacao,'GRD':x.grd,'Data Emissão':x.dataEmissao,'PCF':x.pcf,'Data Recebimento PCF':x.dataPcf,'Status PCF':x.statusPcf,'PCF Respondida':x.pcfRespondida,'Data Resposta':x.dataResposta,'Comentários OPEN':x.open,'Under Review':x.underReview,'Responsável':x.responsavel,'Documento KM':x.documentoKm,'Transmittal KM':x.transmittalKm,'Data KM':x.dataKm}));const wb=XLSX.utils.book_new(),ws=XLSX.utils.json_to_sheet(rows);ws['!autofilter']={ref:ws['!ref']};ws['!cols']=Array(25).fill({wch:20});XLSX.utils.book_append_sheet(wb,ws,'LD Projeto Básico');XLSX.writeFile(wb,'LD_Projeto_Basico_recorte.xlsx')});init();update();
</script></body></html>'''


def build(source: Path = DEFAULT_SOURCE, output: Path = DEFAULT_OUTPUT):
    records = load_records(source)
    meta = {"generated": datetime.now().astimezone().isoformat(), "source": str(source), "records": len(records)}
    html = HTML.replace("__DATA__", json.dumps(records, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/"))
    html = html.replace("__META__", json.dumps(meta, ensure_ascii=False, separators=(",", ":")))
    logo = "data:image/png;base64," + base64.b64encode(DEFAULT_LOGO.read_bytes()).decode("ascii")
    html = html.replace("__LOGO__", logo)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_suffix(output.suffix + ".tmp")
    temp.write_text(html, encoding="utf-8")
    temp.replace(output)
    return {"output": str(output), "records": len(records), "bytes": output.stat().st_size}


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT
    print(json.dumps(build(source, output), ensure_ascii=False))
