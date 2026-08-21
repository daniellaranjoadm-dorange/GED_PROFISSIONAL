"""Gera o dashboard gerencial de Document Control da LD Projeto Básico."""

from __future__ import annotations

import json
import base64
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

try:
    from scripts.build_ld_executive_dashboard import (
        DEFAULT_SOURCE, clean_pcf_status, date_or_text, export_cell, integer, iso,
        load_ld_headers, load_records
    )
except ModuleNotFoundError:  # execução direta a partir da pasta scripts
    from build_ld_executive_dashboard import (
        DEFAULT_SOURCE, clean_pcf_status, date_or_text, export_cell, integer, iso,
        load_ld_headers, load_records
    )


DEFAULT_TEMPLATE = Path(
    r"\\virm-rgr022\FILESERVER\Projetos\05_HANDYMAX\09. Doc Control\3 - LD"
    r"\Dashboard_Doc_Control_LD_Projeto_Bascico"
    r"\Dashboard Doc Control_LD Projeto Bascico.html"
)
DEFAULT_OUTPUT = Path(
    r"\\virm-rgr022\FILESERVER\Projetos\05_HANDYMAX\09. Doc Control\3 - LD"
    r"\Dashboard_Gerencial_Doc_Control_LD_Projeto_Basico.html"
)
DEFAULT_LOGO_ICO = Path(r"D:\GED_PROFISSIONAL\Logo icone de Pasta.ico")
LD_MAX_COLUMN = 61  # A:BI
JSZIP_BROWSER = Path(
    r"C:\Users\daniel.laranjo\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\node\node_modules\jszip\dist\jszip.min.js"
)
PPTXGENJS_BROWSER = Path(
    r"C:\Users\daniel.laranjo\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\node\node_modules\pptxgenjs\dist\pptxgen.min.js"
)


def resolve_template(template: Path) -> Path:
    """Localiza o template mesmo se a pasta receber outro prefixo numérico."""
    if template.is_file():
        return template

    nome_arquivo = template.name
    raiz_ld = template.parent.parent
    candidatos = sorted(
        raiz_ld.glob(f"*Dashboard_Doc_Control_LD_Projeto_Bascico/{nome_arquivo}")
    )
    if candidatos:
        return candidatos[-1]

    raise FileNotFoundError(
        f"Template do dashboard não encontrado: {template}. "
        f"Também foi pesquisado em: {raiz_ld}"
    )


def _safe_json(value) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":")).replace(
        "</", "<\\/"
    )


def load_full_records(source: Path):
    """Le todas as linhas validas da aba, preservando revisoes e repeticoes."""
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    sheet = workbook["LD PROJETO BASICO"]
    records = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=LD_MAX_COLUMN, values_only=True), 2
    ):
        document = str(values[2] or "").strip()
        if not document or document.upper() in {"NOT APPLICABLE", "N/A", "#N/A"}:
            continue
        revision = values[3] if values[3] not in (None, "") else 0
        records.append(
            {
                "documento": document,
                "tipo": str(values[0] or "-").strip(),
                "revisao": str(revision),
                "titulo": str(values[4] or "-").replace("\n", " ").strip(),
                "disciplina": str(values[6] or values[5] or "SEM DISCIPLINA").strip(),
                "especialidade": str(values[14] or "-").strip(),
                "status": str(values[12] or "SEM STATUS").strip(),
                "statusEmissao": str(values[13] or "SEM STATUS").strip(),
                "grd": str(values[22] or "").strip(),
                "dataEmissao": iso(values[23]),
                "pcf": str(values[24] or "").strip(),
                "dataPcf": iso(values[25]),
                "statusPcf": clean_pcf_status(values[26]),
                "pcfRespondida": str(values[27] or "").strip(),
                "dataResposta": iso(values[28]),
                "grdResposta": str(values[29] or "").strip(),
                "responsavel": str(values[21] or "SEM RESPONSÁVEL").strip(),
                "guiaEmissao": str(values[34] or "").strip(),
                "statusDox": str(values[46] or "SEM REGISTRO").strip(),
                "documentoKm": str(values[7] or "").strip(),
                "kmTitle": str(values[8] or "").replace("\n", " ").strip(),
                "transmittalKm": str(values[9] or "").strip(),
                "dataKm": iso(values[10]),
                "documentoKmEmitidoDocTp": str(values[11] or "").strip(),
                "medicaoEmissao": date_or_text(values[47]),
                "medicaoAprovacao": date_or_text(values[48]),
                "casco": str(values[54] or "").strip(),
                "open": integer(values[31]),
                "comentarios": integer(values[30]),
                "underReview": integer(values[32]),
                "statusFinalPcf": clean_pcf_status(values[33]),
                "cronogramaInicio": iso(values[55]),
                "cronogramaTermino": iso(values[56]),
                "bmEmissao": str(values[57] or "SEM INFORMAÇÃO").strip(),
                "bmDataEmissao": date_or_text(values[58]),
                "bmAprovacao": str(values[59] or "SEM INFORMAÇÃO").strip(),
                "bmDataAprovacao": date_or_text(values[60]),
                "linhaFonte": row_number,
                "ldExport": [export_cell(value) for value in values],
            }
        )
    workbook.close()
    return records


def build(
    source: Path = DEFAULT_SOURCE,
    output: Path = DEFAULT_OUTPUT,
    template: Path = DEFAULT_TEMPLATE,
):
    records = load_records(source)
    generated = datetime.now().astimezone().isoformat()
    meta = {
        "generated": generated,
        "source": str(source),
        "records": len(records),
        "mode": "dashboard gerencial consolidado",
        "ldHeaders": load_ld_headers(source),
    }

    template = resolve_template(template)
    html = template.read_text(encoding="utf-8")
    logo_uri = "data:image/x-icon;base64," + base64.b64encode(
        DEFAULT_LOGO_ICO.read_bytes()
    ).decode("ascii")
    html = html.replace(
        '<div class="mark">DC</div>',
        f'<div class="mark"><img src="{logo_uri}" alt="D’OR@NGE"></div>',
        1,
    )
    data_start = html.index("const DATA=")
    runtime_start = html.index("const $=id", data_start)
    runtime = f"const DATA={_safe_json(records)},META={_safe_json(meta)};const FILES={{}};\n"
    html = html[:data_start] + runtime + html[runtime_start:]

    # A interface e os downloads usam a mesma base consolidada.
    test_style = """
<style>
  .filter-grid{grid-template-columns:repeat(5,minmax(0,1fr))!important}
  .v2-filter{position:relative;min-width:0}
  .v2-filter>span{display:block;margin:0 0 6px;color:#8faab7;font-size:10px;
    font-weight:800;letter-spacing:.08em;text-transform:uppercase}
  .v2-multi-btn{width:100%;height:42px;padding:0 12px;border:1px solid #294d64;
    border-radius:10px;background:#071925;color:#eaf7fb;text-align:left;cursor:pointer;
    overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
  .v2-filter.open .v2-multi-btn{border-color:#37c8f4;box-shadow:0 0 0 3px #37c8f422}
  .v2-panel{display:none;position:absolute;z-index:200;top:calc(100% + 6px);left:0;
    width:max-content;min-width:100%;max-width:340px;max-height:300px;overflow:auto;
    padding:7px;border:1px solid #315b73;border-radius:11px;background:#071925;
    box-shadow:0 20px 45px #000b}
  .v2-filter.open .v2-panel{display:block}
  .v2-option{display:flex;align-items:center;gap:8px;padding:7px 8px;border-radius:7px;
    color:#dcebf2;font-size:12px;white-space:nowrap;cursor:pointer}
  .v2-option:hover{background:#123047}.v2-option:first-child{color:#37c8f4;font-weight:800;
    border-bottom:1px solid #294d64;margin-bottom:4px}
  .v2-option input{accent-color:#37c8f4}
  .v2-toolbar{display:flex;flex-wrap:wrap;align-items:center;gap:8px;margin-top:14px}
  .v2-toolbar .control{flex:1 1 380px}.v2-export{height:42px;padding:0 15px;border-radius:10px;
    border:1px solid #315b73;background:#0c2738;color:#dcecf3;font-weight:800;cursor:pointer}
  .v2-export.primary{border-color:#37c8f4;background:#37c8f4;color:#06131e}
  .actions .resource{min-height:68px;padding:10px 15px;border-width:2px;box-shadow:0 12px 30px #0005}
  .actions .resource .copy{min-width:145px}.actions .resource .copy small{color:#9adff3}
  .actions .resource .copy strong{font-size:14px}.actions .resource:after{content:"BASE GERAL";
    align-self:flex-start;padding:4px 6px;border-radius:999px;background:#173f54;color:#7ddcff;
    font-size:8px;font-weight:950;letter-spacing:.08em;white-space:nowrap}
  .actions .resource.ppt:after{content:"GERAL";background:#123c34;color:#65ead0}
  .actions .utility.primary{min-height:54px;padding:11px 17px;border-width:2px;box-shadow:0 10px 24px #0780b044}
  .mark{overflow:hidden;padding:0;background:#102b3d!important}.mark img{width:100%;height:100%;object-fit:cover;display:block}
  .v2-summary{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px;margin-top:14px}
  .v2-summary>div{padding:11px 13px;border:1px solid #234359;border-radius:11px;background:#071925}
  .v2-summary strong{display:block;color:#fff;font-size:19px}.v2-summary span{color:#8faab7;font-size:10px}
  .doc-link{color:#7ddcff;text-decoration:none;text-underline-offset:3px}.doc-link:hover{color:#fff;text-decoration:underline}.doc-link span{font-size:10px;color:#38bdf8}
  .kpis{grid-template-columns:repeat(4,minmax(0,1fr))!important}
  .v2-start{display:flex;align-items:center;justify-content:space-between;gap:16px;
    margin:-2px 0 18px;padding:14px 16px;border:1px solid rgba(55,200,244,.38);
    border-left:4px solid #37c8f4;border-radius:12px;background:linear-gradient(90deg,#0b2a3c,#081c29)}
  .v2-start strong{display:block;color:#fff;font-size:14px}.v2-start span{color:#9fc0cf;font-size:11px}
  .v2-start b{flex:0 0 auto;padding:7px 10px;border-radius:999px;background:#37c8f4;
    color:#06131e;font-size:10px;letter-spacing:.08em;text-transform:uppercase}
  .nav button[data-view="base"]{border-color:#37c8f4!important;background:linear-gradient(135deg,#173f54,#0e2d3e)!important}
  .nav button[data-view="base"] strong:after{content:" · COMECE AQUI";color:#37c8f4;font-size:9px;letter-spacing:.08em}
  @media(max-width:1100px){.filter-grid{grid-template-columns:repeat(3,minmax(0,1fr))!important}.kpis{grid-template-columns:repeat(2,minmax(0,1fr))!important}}
  @media(max-width:700px){.filter-grid{grid-template-columns:1fr!important}.v2-summary{grid-template-columns:repeat(2,1fr)}.kpis{grid-template-columns:1fr!important}}
</style>
"""
    html = html.replace("</head>", test_style + "</head>", 1)
    enhancement_script = r'''
<script>
(()=>{
  const fields=[
    ['tipo','Tipo'],['disciplina','Disciplina'],['status','Status geral'],
    ['statusEmissao','Emissão'],['_prazo','Prazo do cronograma'],
    ['medicaoEmissao','Medição emissão'],['medicaoAprovacao','Medição aprovação'],
    ['statusPcf','Status PCF'],['responsavel','Responsável'],['casco','Casco'],
    ['bmEmissao','BM - Emissão'],['bmAprovacao','BM - Aprovação']
  ];
  const selectedSets={},fmtV2=new Intl.NumberFormat('pt-BR');
  const valueOf=(r,k)=>k==='_prazo'?deadline(r):String(r[k]||'-');
  const values=k=>[...new Set(DATA.map(r=>valueOf(r,k)))].sort((a,b)=>String(a).localeCompare(String(b),'pt-BR'));
  const grid=document.querySelector('.filter-grid');
  if(!grid)return;
  const section=grid.closest('section');
  const heading=section?.querySelector('h2,h3');if(heading)heading.textContent='Central de documentos e filtros';
  const intro=heading?.parentElement?.querySelector('p');if(intro)intro.textContent='Consulte a carteira, aplique filtros e escolha claramente entre a visão gerencial ou a LD completa.';
  const start=document.createElement('div');start.className='v2-start';start.innerHTML='<div><strong>Precisa localizar, analisar ou entregar documentos?</strong><span>Comece nesta área. Os filtros atualizam a tabela e o resumo abaixo.</span></div><b>Área principal</b>';grid.before(start);
  grid.innerHTML=fields.map(([k,label])=>`<div class="v2-filter" data-field="${k}"><span>${label}</span><button type="button" class="v2-multi-btn">Todos</button><div class="v2-panel"></div></div>`).join('');
  const toolbar=document.createElement('div');toolbar.className='v2-toolbar';toolbar.innerHTML=`
    <input class="control" id="v2Search" placeholder="Buscar documento, título, GRD, PCF ou KM...">
    <button class="v2-export" id="v2Clear">Limpar filtros</button>
    <button class="v2-export primary" id="v2ExcelOfficial">Excel do recorte filtrado</button>
    <button class="v2-export" id="v2Pptx">PPTX do recorte filtrado</button>
    <button class="v2-export" id="v2Pdf">PDF do recorte filtrado</button>`;
  grid.after(toolbar);
  const summary=document.createElement('div');summary.className='v2-summary';summary.id='v2Summary';toolbar.after(summary);
  const esc2=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
  const safeLink=u=>/^https?:\/\//i.test(String(u||''))?String(u):'';
  const linked=(label,url,title)=>{const safe=safeLink(url);return safe?`<a class="doc-link" href="${esc2(safe)}" target="_blank" rel="noopener noreferrer" title="${esc2(title)}">${esc2(label)} <span aria-hidden="true">&#8599;</span></a>`:esc2(label||'-')};
  const legacyRenderTable=renderTable;
  renderTable=function(){legacyRenderTable();const rows=filtered().sort((a,b)=>priority(b)-priority(a)).slice(0,250),trs=[...document.querySelectorAll('#rows tr')];trs.forEach((tr,index)=>{const r=rows[index];if(!r)return;const cells=tr.children;if(cells[0])cells[0].innerHTML='<strong>'+linked(r.documento,r.linkDox,'Abrir documento no DOX')+'</strong>';if(cells[8])cells[8].innerHTML=linked(r.pcf||'-',r.linkPcf,'Abrir ultima versao da PCF no DOX')})};
  function setupFilter(box,k){
    const vals=values(k);selectedSets[k]=new Set(vals);const panel=box.querySelector('.v2-panel'),btn=box.querySelector('button');
    panel.innerHTML=`<label class="v2-option"><input type="checkbox" data-all checked> Selecionar tudo</label>`+vals.map(v=>`<label class="v2-option"><input type="checkbox" data-value="${esc2(v)}" checked> ${esc2(v)}</label>`).join('');
    const all=panel.querySelector('[data-all]'),items=[...panel.querySelectorAll('[data-value]')];
    const sync=()=>{selectedSets[k]=new Set(items.filter(x=>x.checked).map(x=>x.dataset.value));all.checked=items.every(x=>x.checked);all.indeterminate=!all.checked&&items.some(x=>x.checked);btn.textContent=all.checked?'Todos':selectedSets[k].size?`${selectedSets[k].size} selecionados`:'Nenhum';refreshV2()};
    btn.onclick=e=>{e.stopPropagation();document.querySelectorAll('.v2-filter.open').forEach(x=>x!==box&&x.classList.remove('open'));box.classList.toggle('open')};
    all.onchange=()=>{items.forEach(x=>x.checked=all.checked);sync()};items.forEach(x=>x.onchange=sync);
  }
  [...grid.querySelectorAll('.v2-filter')].forEach(box=>setupFilter(box,box.dataset.field));
  filtered=function(){const q=(document.getElementById('v2Search')?.value||'').trim().toLocaleLowerCase('pt-BR');return DATA.filter(r=>fields.every(([k])=>selectedSets[k].has(valueOf(r,k)))&&(!q||[r.documento,r.titulo,r.grd,r.pcf,r.documentoKm,r.transmittalKm].some(v=>String(v||'').toLocaleLowerCase('pt-BR').includes(q))))};
  function refreshV2(){renderTable();const rows=filtered(),emitted=rows.filter(r=>r.statusEmissao==='Emitido').length,pending=rows.filter(r=>r.statusEmissao!=='Emitido').length,overdue=rows.filter(r=>deadline(r)==='Vencido').length;summary.innerHTML=`<div><strong>${fmtV2.format(rows.length)}</strong><span>DOCUMENTOS CONSOLIDADOS</span></div><div><strong>${fmtV2.format(emitted)}</strong><span>EMITIDOS</span></div><div><strong>${fmtV2.format(pending)}</strong><span>PENDENTES</span></div><div><strong>${fmtV2.format(overdue)}</strong><span>VENCIDOS</span></div>`}
  document.getElementById('v2Search').oninput=refreshV2;document.addEventListener('click',()=>document.querySelectorAll('.v2-filter.open').forEach(x=>x.classList.remove('open')));
  document.getElementById('v2Clear').onclick=()=>{document.getElementById('v2Search').value='';grid.querySelectorAll('input[type=checkbox]').forEach(x=>x.checked=true);fields.forEach(([k])=>selectedSets[k]=new Set(values(k)));grid.querySelectorAll('.v2-multi-btn').forEach(x=>x.textContent='Todos');refreshV2()};
  const columns=[['Linha fonte','linhaFonte'],['Documento','documento'],['Revisão','revisao'],['Título','titulo'],['Tipo','tipo'],['Disciplina','disciplina'],['Especialidade','especialidade'],['Status','status'],['Status Emissão','statusEmissao'],['Início cronograma','cronogramaInicio'],['Término cronograma','cronogramaTermino'],['Prazo','_prazo'],['Medição emissão','medicaoEmissao'],['Medição aprovação','medicaoAprovacao'],['GRD','grd'],['Data emissão','dataEmissao'],['PCF','pcf'],['Data PCF','dataPcf'],['Status PCF','statusPcf'],['PCF respondida','pcfRespondida'],['Data resposta','dataResposta'],['GRD resposta','grdResposta'],['Comentários','comentarios'],['OPEN','open'],['Under review','underReview'],['Responsável','responsavel'],['Status DOX','statusDox'],['Guia emissão','guiaEmissao'],['Casco','casco'],['BM - Emissão','bmEmissao'],['BM - Data Emissão','bmDataEmissao'],['BM - Aprovação','bmAprovacao'],['BM - Data Aprovação','bmDataAprovacao'],['Documento KM','documentoKm'],['Transmittal KM','transmittalKm'],['Data KM','dataKm']];
  const cell=(r,k)=>k==='_prazo'?deadline(r):(r[k]??'');
  const download=(blob,name)=>{const u=URL.createObjectURL(blob),a=document.createElement('a');a.href=u;a.download=name;a.click();setTimeout(()=>URL.revokeObjectURL(u),1200)};
  const ensureRows=()=>{const rows=filtered();if(!rows.length)alert('O recorte atual não possui documentos para exportar.');return rows};
  const excelColumn=n=>{let out='';while(n){n--;out=String.fromCharCode(65+n%26)+out;n=Math.floor(n/26)}return out};
  const xmlEscape=v=>String(v??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&apos;'}[c]));
  const xlsxCell=(value,address,style=0)=>{if(typeof value==='number'&&Number.isFinite(value))return `<c r="${address}" s="${style}"><v>${value}</v></c>`;if(typeof value==='boolean')return `<c r="${address}" s="${style}" t="b"><v>${value?1:0}</v></c>`;return `<c r="${address}" s="${style}" t="inlineStr"><is><t xml:space="preserve">${xmlEscape(value)}</t></is></c>`};
  async function exportFilteredExcel(){const rows=ensureRows();if(!rows.length)return;if(typeof JSZip==='undefined'){alert('O gerador de Excel não foi carregado. Atualize o dashboard.');return}const headers=META.ldHeaders||[],matrix=[headers,...rows.map(r=>headers.map((_,i)=>r.ldExport?.[i]??''))],lastCol=excelColumn(headers.length),lastRow=matrix.length,zip=new JSZip(),sheetRows=matrix.map((row,ri)=>`<row r="${ri+1}"${ri===0?' ht="30" customHeight="1"':''}>${row.map((v,ci)=>xlsxCell(v,`${excelColumn(ci+1)}${ri+1}`,ri===0?1:0)).join('')}</row>`).join(''),cols=headers.map((h,i)=>{const width=/TITLE|T.TULO/i.test(h)?42:/DISCIPLIN/i.test(h)?32:/DOCUMENTO|TRANSMITTAL|GRD|PCF/i.test(h)?26:18;return `<col min="${i+1}" max="${i+1}" width="${width}" customWidth="1"/>`}).join('');zip.file('[Content_Types].xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>');zip.folder('_rels').file('.rels','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>');zip.folder('xl').file('workbook.xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="LD Recorte Filtrado" sheetId="1" r:id="rId1"/></sheets></workbook>');zip.folder('xl').folder('_rels').file('workbook.xml.rels','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>');zip.folder('xl').file('styles.xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="2"><font><sz val="10"/><name val="Arial"/></font><font><b/><color rgb="FFFFFFFF"/><sz val="10"/><name val="Arial"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF173F54"/><bgColor indexed="64"/></patternFill></fill></fills><borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf></cellXfs></styleSheet>');zip.folder('xl').folder('worksheets').file('sheet1.xml',`<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:${lastCol}${lastRow}"/><sheetViews><sheetView workbookViewId="0"><pane xSplit="2" ySplit="1" topLeftCell="C2" activePane="bottomRight" state="frozen"/></sheetView></sheetViews><cols>${cols}</cols><sheetData>${sheetRows}</sheetData><autoFilter ref="A1:${lastCol}${lastRow}"/></worksheet>`);const blob=await zip.generateAsync({type:'blob',compression:'DEFLATE'});download(blob,`LD_Projeto_Basico_Recorte_${rows.length}_documentos.xlsx`)}
  async function exportFilteredPptx(){const rows=ensureRows();if(!rows.length)return;if(typeof PptxGenJS==='undefined'){alert('O gerador de PowerPoint não foi carregado. Atualize o dashboard.');return}const pptx=new PptxGenJS();pptx.layout='LAYOUT_WIDE';pptx.author='D’OR@NGE · Document Control';pptx.subject='Recorte filtrado da LD Projeto Básico';pptx.title='LD Projeto Básico · Recorte filtrado';pptx.company='Consórcio Marenova';const bg='06131E',white='F4FBFE',muted='9BB6C3',cyan='37C8F4',mint='1DD6B5',amber='FFAD32',red='FF6670';let s=pptx.addSlide();s.background={color:bg};s.addText('D’OR@NGE · DOCUMENT CONTROL',{x:.65,y:.45,w:6,h:.3,fontSize:11,bold:true,color:amber});s.addText('LD Projeto Básico',{x:.65,y:1.5,w:8.5,h:.65,fontSize:34,bold:true,color:white});s.addText('Apresentação gerencial do recorte filtrado',{x:.65,y:2.25,w:9,h:.4,fontSize:20,color:cyan});s.addText(`${rows.length} documentos · gerado em ${new Date().toLocaleString('pt-BR')}`,{x:.65,y:3.05,w:8,h:.3,fontSize:13,color:muted});const emitted=rows.filter(r=>r.statusEmissao==='Emitido').length,overdue=rows.filter(r=>deadline(r)==='Vencido').length,open=rows.reduce((n,r)=>n+(Number(r.open)||0),0);s=pptx.addSlide();s.background={color:bg};s.addText('Resumo do recorte',{x:.6,y:.4,w:8,h:.45,fontSize:26,bold:true,color:white});[['Documentos',rows.length,cyan],['Emitidos',emitted,mint],['Pendentes',rows.length-emitted,amber],['Vencidos',overdue,red],['Comentários OPEN',open,amber]].forEach((c,i)=>{const x=.6+i*2.5;s.addShape(pptx.ShapeType.roundRect,{x,y:1.35,w:2.2,h:1.25,rectRadius:.08,fill:{color:'0D2233'},line:{color:c[2],width:1.5}});s.addText(c[0],{x:x+.15,y:1.55,w:1.9,h:.25,fontSize:10,bold:true,color:muted});s.addText(String(c[1]),{x:x+.15,y:1.92,w:1.9,h:.4,fontSize:25,bold:true,color:c[2]})});s.addText('Os indicadores e a lista abaixo refletem exclusivamente os filtros aplicados na Central de Documentos.',{x:.6,y:3.15,w:11.8,h:.45,fontSize:15,color:muted});s=pptx.addSlide();s.background={color:bg};s.addText('Documentos prioritários do recorte',{x:.6,y:.35,w:9,h:.4,fontSize:24,bold:true,color:white});const tableRows=[['Documento','Rev.','Disciplina','Status','Emissão','Prazo'],...rows.slice().sort((a,b)=>priority(b)-priority(a)).slice(0,18).map(r=>[r.documento,r.revisao,r.disciplina,r.status,r.statusEmissao,deadline(r)])];s.addTable(tableRows,{x:.55,y:1,w:12.2,h:5.8,border:{type:'solid',color:'294D64',pt:1},fill:'0D2233',color:white,fontSize:8,margin:.05,autoFit:false,colW:[3.1,.55,2.65,2.1,1.45,1.25],bold:false,rowH:.28});await pptx.writeFile({fileName:`LD_Projeto_Basico_Recorte_${rows.length}_documentos.pptx`})}
  function exportFilteredPdf(){const rows=ensureRows();if(!rows.length)return;const popup=window.open('','_blank');if(!popup){alert('Permita pop-ups para gerar o PDF do recorte.');return}const emitted=rows.filter(r=>r.statusEmissao==='Emitido').length,overdue=rows.filter(r=>deadline(r)==='Vencido').length,reportColumns=[['Documento','documento'],['Rev.','revisao'],['Título','titulo'],['Disciplina','disciplina'],['Status','status'],['Emissão','statusEmissao'],['Prazo','_prazo'],['GRD','grd'],['PCF','pcf'],['Status PCF','statusPcf'],['Responsável','responsavel'],['BM Emissão','bmEmissao'],['BM Aprovação','bmAprovacao']];popup.document.write(`<!doctype html><html><head><meta charset="utf-8"><title>Relatório filtrado · LD Projeto Básico</title><style>@page{size:A4 landscape;margin:10mm}body{font:11px Arial;color:#17212b}h1{margin:0;color:#123b50}.meta{color:#557080;margin:5px 0 15px}.cards{display:flex;gap:8px;margin:12px 0}.card{border:1px solid #9fb5c0;padding:8px 12px;min-width:120px}.card b{display:block;font-size:20px;color:#123b50}table{width:100%;border-collapse:collapse;font-size:7px}th{background:#123b50;color:#fff}th,td{border:1px solid #ccd8de;padding:3px;text-align:left}tr:nth-child(even){background:#eef4f6}</style></head><body><h1>LD Projeto Básico · Relatório do recorte</h1><div class="meta">Gerado em ${new Date().toLocaleString('pt-BR')} · filtros da Central de Documentos</div><div class="cards"><div class="card"><b>${rows.length}</b>Documentos</div><div class="card"><b>${emitted}</b>Emitidos</div><div class="card"><b>${rows.length-emitted}</b>Pendentes</div><div class="card"><b>${overdue}</b>Vencidos</div></div><table><thead><tr>${reportColumns.map(c=>`<th>${esc2(c[0])}</th>`).join('')}</tr></thead><tbody>${rows.map(r=>`<tr>${reportColumns.map(c=>`<td>${esc2(cell(r,c[1]))}</td>`).join('')}</tr>`).join('')}</tbody></table><script>onload=()=>setTimeout(()=>print(),250)<\/script></body></html>`);popup.document.close()}
  document.getElementById('v2ExcelOfficial').onclick=exportFilteredExcel;
  document.getElementById('v2Pptx').onclick=exportFilteredPptx;
  document.getElementById('v2Pdf').onclick=exportFilteredPdf;refreshV2();
  const emittedCount=DATA.filter(r=>r.statusEmissao==='Emitido').length;
  const approvedNoComments=DATA.filter(r=>String(r.status||'').toLocaleLowerCase('pt-BR')==='aprovado sem comentários').length;
  MEM.emittedCount={title:'Documentos emitidos',subtitle:'Quantidade absoluta de documentos formalmente emitidos',result:fmt(emittedCount),color:'#1dd6b5',formula:'Contagem dos documentos em que Status Emissão = “Emitido”.',values:[['Documentos emitidos',fmt(emittedCount)],['Total consolidado',fmt(DATA.length)],['Percentual da carteira',pct(emittedCount,DATA.length)]],reading:'Volume efetivamente emitido na última revisão válida de cada documento.',source:'Campo Status Emissão da aba LD PROJETO BASICO.'};
  MEM.approvedNoComments={title:'Aprovados sem comentários',subtitle:'Documentos aprovados sem ressalvas registradas',result:fmt(approvedNoComments),color:'#48d7a8',formula:'Contagem dos documentos em que Status = “Aprovado sem Comentários”.',values:[['Aprovados sem comentários',fmt(approvedNoComments)],['Total consolidado',fmt(DATA.length)],['Percentual da carteira',pct(approvedNoComments,DATA.length)]],reading:'Representa documentos aprovados sem necessidade de tratamento adicional de comentários na revisão consolidada.',source:'Campo Status da aba LD PROJETO BASICO.'};
  const kpiGrid=document.getElementById('kpiGrid');if(kpiGrid){kpiGrid.insertAdjacentHTML('beforeend',`<article class="kpi glass" style="--a:#1dd6b5" data-memory="emittedCount" tabindex="0"><div class="label">Documentos emitidos</div><div class="num">${fmt(emittedCount)}</div><div class="sub">quantidade absoluta emitida</div><div class="hint"><i>+</i>Ver memória</div></article><article class="kpi glass" style="--a:#48d7a8" data-memory="approvedNoComments" tabindex="0"><div class="label">Aprovados sem comentários</div><div class="num">${fmt(approvedNoComments)}</div><div class="sub">aprovação sem ressalvas</div><div class="hint"><i>+</i>Ver memória</div></article>`)}
  const baseNav=document.querySelector('.nav button[data-view="base"]');if(baseNav){baseNav.querySelector('.nav-index').textContent='01';const strong=baseNav.querySelector('strong');if(strong)strong.textContent='Central de documentos';baseNav.querySelector('.nav-copy').lastChild.textContent=' Filtros, auditoria e exportações';document.querySelector('.nav').prepend(baseNav);baseNav.click()}
  const legacyActions=[document.getElementById('clearFilters'),document.getElementById('exportCsv')].filter(Boolean);legacyActions.forEach(x=>x.style.display='none');
  document.querySelectorAll('.resource').forEach(b=>{if(/XLSM/i.test(b.textContent)){b.querySelector('.fileicon').textContent='XLSX';b.querySelector('small').textContent='Base geral consolidada';b.querySelector('strong').textContent='Excel completo';b.setAttribute('onclick',"downloadFile('xlsx')")}if(/PPTX/i.test(b.textContent)){b.querySelector('small').textContent='Apresentação geral';b.querySelector('strong').textContent='PowerPoint completo'}});
  const headerBase=[...document.querySelectorAll('header button')].find(b=>/Base detalhada/i.test(b.textContent));if(headerBase){headerBase.textContent='Abrir Central de Documentos';headerBase.classList.add('primary')}
  const headerPdf=[...document.querySelectorAll('header button')].find(b=>/Exportar PDF/i.test(b.textContent));if(headerPdf)headerPdf.textContent='PDF geral completo';
})();
</script>
'''
    pptxgenjs = "\n".join(
        path.read_text(encoding="utf-8") for path in (JSZIP_BROWSER, PPTXGENJS_BROWSER)
    ).replace("</script", "<\\/script")
    html = html.replace("</head>", f"<script>{pptxgenjs}</script></head>", 1)
    html = html.replace(
        "</body>",
        enhancement_script + "</body>",
        1,
    )
    html = html.replace(
        "<title>Dashboard Doc Control_LD Projeto Bascico</title>",
        "<title>Dashboard Gerencial Doc Control · LD Projeto Básico</title>",
        1,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    export_dir = output.parent / "Dashboard_Doc_Control_Arquivos"
    build_dir = Path(tempfile.gettempdir()) / "ged_ld_dashboard_build"
    preview_dir = build_dir / "previews"
    export_dir.mkdir(parents=True, exist_ok=True)
    build_dir.mkdir(parents=True, exist_ok=True)
    payload_path = build_dir / "dashboard_export_data.json"
    payload_path.write_text(_safe_json({"records": records, "meta": meta}), encoding="utf-8")
    exporter = build_dir / "build_ld_management_exports.mjs"
    shutil.copy2(Path(__file__).with_name("build_ld_management_exports.mjs"), exporter)
    node_exe = Path(r"C:\Users\daniel.laranjo\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
    node_modules_link = build_dir / "node_modules"
    node_modules_source = Path(r"C:\Users\daniel.laranjo\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules")
    if not node_modules_link.exists():
        subprocess.run(
            ["cmd.exe", "/c", "mklink", "/J", str(node_modules_link), str(node_modules_source)],
            check=True,
            capture_output=True,
            text=True,
        )
    xlsx_path = export_dir / "LD_Projeto_Basico_Consolidada.xlsx"
    pptx_path = export_dir / "LD_Projeto_Basico_Apresentacao_Gerencial.pptx"
    subprocess.run([str(node_exe), str(exporter), str(payload_path), str(xlsx_path), str(pptx_path), str(preview_dir)], check=True)
    files = {
        "xlsx": {"name": xlsx_path.name, "type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "base64": base64.b64encode(xlsx_path.read_bytes()).decode("ascii")},
        "pptx": {"name": pptx_path.name, "type": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "base64": base64.b64encode(pptx_path.read_bytes()).decode("ascii")},
    }
    html = html.replace("const FILES={};", f"const FILES={_safe_json(files)};")
    temporary = output.with_suffix(output.suffix + ".tmp")
    temporary.write_text(html, encoding="utf-8")
    temporary.replace(output)
    return {
        "output": str(output),
        "records": len(records),
        "xlsx": str(xlsx_path),
        "pptx": str(pptx_path),
        "bytes": output.stat().st_size,
        "generated": generated,
        "published": False,
    }


if __name__ == "__main__":
    source_arg = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    output_arg = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT
    print(json.dumps(build(source_arg, output_arg, DEFAULT_TEMPLATE), ensure_ascii=False))
