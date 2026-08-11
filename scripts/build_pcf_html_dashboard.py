import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\daniel.laranjo\Downloads\controle_respostas_pcf_LD_PROJETO_BASICO (2).xlsx")
OUTPUT = Path(
    r"C:\Users\daniel.laranjo\.codex\visualizations\2026\08\03\019fc763-53de-7740-999c-b52eea40f680"
    r"\outputs\pcf_html_dashboard\dashboard_pcf_diretoria.html"
)


def clean_status(value):
    value = str(value or "-").strip().upper()
    if value in {"NOT RELESED", "NOT RELEASED"}:
        return "NOT RELEASED"
    return value or "-"


def serializable(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return None
    return value


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Dados PCFs"]
headers = [cell.value for cell in next(sheet.iter_rows())]
records = []
for values in sheet.iter_rows(min_row=2, values_only=True):
    if not values[0]:
        continue
    row = {header: serializable(value) for header, value in zip(headers[:24], values[:24])}
    row["Status PCF"] = clean_status(row.get("Status PCF"))
    for field in ("Dias sem Resposta", "Dias de Atraso", "Comentarios", "Open", "Under Review", "Closed Calculado"):
        row[field] = int(row.get(field) or 0)
    records.append(row)
workbook.close()

data_json = json.dumps(records, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")

html = r'''<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>PCF Command Center | LD Projeto Básico</title>
  <script src="https://cdn.tailwindcss.com"></script>
  <script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.7/dist/chart.umd.min.js"></script>
  <script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
  <script>
    tailwind.config={theme:{extend:{fontFamily:{sans:['Inter','Segoe UI','sans-serif']},colors:{ink:'#071521',panel:'#0d2233',line:'#1d3b50',cyan:'#39c6f4',amber:'#f5a524',danger:'#ff5c5c',mint:'#20d6b5'}}}}
  </script>
  <style>
    :root{color-scheme:dark}*{box-sizing:border-box}body{background:radial-gradient(circle at 10% 0%,#12334a 0,#071521 34%,#041019 100%);min-height:100vh}
    .glass{background:linear-gradient(145deg,rgba(13,34,51,.94),rgba(7,21,33,.96));border:1px solid rgba(99,179,225,.18);box-shadow:0 18px 45px rgba(0,0,0,.24)}
    .kpi{position:relative;overflow:hidden}.kpi:before{content:"";position:absolute;inset:0 auto 0 0;width:4px;background:var(--accent)}
    .filter-control{width:100%;border:1px solid #29485c;background:#071925;color:#e8f4fb;border-radius:.7rem;padding:.68rem .78rem;outline:none}.filter-control:focus{border-color:#39c6f4;box-shadow:0 0 0 3px rgba(57,198,244,.12)}
    .chart-box{height:320px}.chart-box-lg{height:360px}.thin-scroll::-webkit-scrollbar{height:8px;width:8px}.thin-scroll::-webkit-scrollbar-thumb{background:#29485c;border-radius:99px}
    .pulse-dot{box-shadow:0 0 0 0 rgba(32,214,181,.6);animation:pulse 2s infinite}@keyframes pulse{70%{box-shadow:0 0 0 8px rgba(32,214,181,0)}100%{box-shadow:0 0 0 0 rgba(32,214,181,0)}}
    @media print{body{background:#fff;color:#111}.glass{box-shadow:none;background:#fff;border:1px solid #ddd}.no-print{display:none!important}.chart-box,.chart-box-lg{height:280px}}
  </style>
</head>
<body class="font-sans text-slate-100 antialiased">
  <main class="mx-auto max-w-[1600px] px-4 py-5 md:px-7 lg:px-9">
    <header class="mb-5 flex flex-col gap-4 lg:flex-row lg:items-end lg:justify-between">
      <div>
        <div class="mb-2 flex items-center gap-2 text-xs font-extrabold uppercase tracking-[.18em] text-cyan"><span class="pulse-dot h-2 w-2 rounded-full bg-mint"></span> Naval Engineering Intelligence</div>
        <h1 class="text-3xl font-black tracking-tight md:text-4xl">PCF Command Center</h1>
        <p class="mt-2 max-w-3xl text-sm text-slate-400">Controle executivo de respostas · LD Projeto Básico · SLA de 15 dias úteis após o recebimento</p>
      </div>
      <div class="flex flex-wrap items-center gap-2 text-xs text-slate-400"><span id="sourceName" class="rounded-full border border-line bg-panel px-3 py-2">Base incorporada</span><span id="dataBase" class="rounded-full border border-line bg-panel px-3 py-2"></span><span id="activeCount" class="rounded-full border border-cyan/30 bg-cyan/10 px-3 py-2 font-bold text-cyan"></span></div>
    </header>

    <section class="glass no-print mb-5 rounded-2xl p-4 md:p-5" aria-label="Filtros interativos">
      <div class="mb-4 flex items-center justify-between gap-3"><div><h2 class="font-extrabold">Filtros interativos</h2><p class="text-xs text-slate-400">Todos os indicadores, gráficos e a carteira crítica respondem instantaneamente.</p></div><button id="clearFilters" class="rounded-lg border border-line px-3 py-2 text-xs font-bold text-slate-300 transition hover:border-cyan hover:text-cyan">Limpar filtros</button></div>
      <div class="grid gap-3 sm:grid-cols-2 lg:grid-cols-4 xl:grid-cols-7">
        <label class="text-xs font-bold text-slate-400">Tipo<select id="tipoFilter" class="filter-control mt-1"></select></label>
        <label class="text-xs font-bold text-slate-400">Situação<select id="situacaoFilter" class="filter-control mt-1"></select></label>
        <label class="text-xs font-bold text-slate-400">Status PCF<select id="statusFilter" class="filter-control mt-1"></select></label>
        <label class="text-xs font-bold text-slate-400">Disciplina<select id="disciplinaFilter" class="filter-control mt-1"></select></label>
        <label class="text-xs font-bold text-slate-400">Responsável<select id="responsavelFilter" class="filter-control mt-1"></select></label>
        <label class="text-xs font-bold text-slate-400">Recebido a partir de<input id="dateFrom" type="date" class="filter-control mt-1"></label>
        <label class="text-xs font-bold text-slate-400">Recebido até<input id="dateTo" type="date" class="filter-control mt-1"></label>
      </div>
      <div class="mt-3 grid gap-3 lg:grid-cols-[1fr_auto]">
        <label class="text-xs font-bold text-slate-400">Buscar documento, título, PCF ou GRD<input id="searchFilter" class="filter-control mt-1" placeholder="Ex.: 5137, máquinas, GRD-234..."></label>
        <div class="mt-5 flex flex-wrap gap-2"><input id="excelUpload" type="file" accept=".xlsx,.xls,.xlsm" class="hidden"><button id="loadExcel" class="rounded-xl border border-cyan/40 bg-cyan/10 px-5 py-3 text-sm font-black text-cyan transition hover:bg-cyan/20">Carregar Excel atualizado</button><button id="exportExcel" class="rounded-xl bg-cyan px-5 py-3 text-sm font-black text-ink transition hover:brightness-110">Exportar recorte Excel</button></div>
      </div>
    </section>

    <section class="mb-5 grid gap-4 sm:grid-cols-2 xl:grid-cols-4" aria-label="Indicadores principais">
      <article class="glass kpi rounded-2xl p-5" style="--accent:#39c6f4"><p class="text-xs font-black uppercase tracking-wider text-cyan">PCFs monitoradas</p><div id="kpiTotal" class="mt-3 text-4xl font-black">—</div><p id="kpiTotalSub" class="mt-2 text-xs text-slate-400"></p></article>
      <article class="glass kpi rounded-2xl p-5" style="--accent:#ff5c5c"><p class="text-xs font-black uppercase tracking-wider text-danger">Vencidas</p><div id="kpiOverdue" class="mt-3 text-4xl font-black">—</div><p id="kpiOverdueSub" class="mt-2 text-xs text-slate-400"></p></article>
      <article class="glass kpi rounded-2xl p-5" style="--accent:#f5a524"><p class="text-xs font-black uppercase tracking-wider text-amber">Comentários OPEN</p><div id="kpiOpen" class="mt-3 text-4xl font-black">—</div><p id="kpiOpenSub" class="mt-2 text-xs text-slate-400"></p></article>
      <article class="glass kpi rounded-2xl p-5" style="--accent:#20d6b5"><p class="text-xs font-black uppercase tracking-wider text-mint">Taxa de resposta</p><div id="kpiResponse" class="mt-3 text-4xl font-black">—</div><p id="kpiResponseSub" class="mt-2 text-xs text-slate-400"></p></article>
    </section>

    <section id="executiveInsight" class="glass mb-5 rounded-2xl border-l-4 border-l-amber p-5"></section>

    <section class="grid gap-5 xl:grid-cols-2">
      <article class="glass rounded-2xl p-5"><div class="mb-4"><p class="text-xs font-black uppercase tracking-wider text-cyan">Carteira</p><h2 class="text-lg font-extrabold">Situação das PCFs</h2></div><div class="chart-box"><canvas id="situationChart"></canvas></div></article>
      <article class="glass rounded-2xl p-5"><div class="mb-4"><p class="text-xs font-black uppercase tracking-wider text-danger">Exposição</p><h2 class="text-lg font-extrabold">Aging do atraso em dias úteis</h2></div><div class="chart-box"><canvas id="agingChart"></canvas></div></article>
      <article class="glass rounded-2xl p-5"><div class="mb-4"><p class="text-xs font-black uppercase tracking-wider text-amber">Carga técnica</p><h2 class="text-lg font-extrabold">Comentários OPEN por disciplina</h2></div><div class="chart-box-lg"><canvas id="disciplineChart"></canvas></div></article>
      <article class="glass rounded-2xl p-5"><div class="mb-4"><p class="text-xs font-black uppercase tracking-wider text-mint">Fluxo</p><h2 class="text-lg font-extrabold">PCFs recebidas por mês</h2></div><div class="chart-box-lg"><canvas id="timelineChart"></canvas></div></article>
    </section>

    <section class="glass mt-5 rounded-2xl p-5">
      <div class="mb-4 flex flex-col gap-2 md:flex-row md:items-end md:justify-between"><div><p class="text-xs font-black uppercase tracking-wider text-danger">Prioridade gerencial</p><h2 class="text-xl font-extrabold">Carteira crítica do recorte</h2><p class="text-xs text-slate-400">Ordenada por atraso e comentários OPEN.</p></div><div id="tableCount" class="text-xs font-bold text-slate-400"></div></div>
      <div class="thin-scroll overflow-x-auto"><table class="w-full min-w-[1050px] text-left text-sm"><thead class="border-b border-line text-[11px] uppercase tracking-wider text-slate-400"><tr><th class="px-3 py-3">Documento / título</th><th class="px-3 py-3">Tipo</th><th class="px-3 py-3">Disciplina</th><th class="px-3 py-3">Situação</th><th class="px-3 py-3 text-right">Atraso</th><th class="px-3 py-3 text-right">Open</th><th class="px-3 py-3">Status PCF</th><th class="px-3 py-3">Responsável</th></tr></thead><tbody id="criticalTable" class="divide-y divide-line/60"></tbody></table></div>
    </section>
    <footer class="py-6 text-center text-xs text-slate-500">D'OR@NGE GED Enterprise · Dashboard local gerado a partir do controle de respostas PCF</footer>
  </main>

<script>
let DATA=__DATA__;
const $=id=>document.getElementById(id);
const fmt=new Intl.NumberFormat('pt-BR');
const pct=v=>`${(v||0).toLocaleString('pt-BR',{maximumFractionDigits:1})}%`;
const esc=s=>String(s??'').replace(/[&<>'"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;',"'":'&#39;','"':'&quot;'}[c]));
const colors={cyan:'#39c6f4',danger:'#ff5c5c',amber:'#f5a524',mint:'#20d6b5',purple:'#9b87f5',slate:'#48657a'};
let charts={}; let filtered=[];
Chart.defaults.color='#a9bdca';Chart.defaults.borderColor='rgba(93,139,166,.18)';Chart.defaults.font.family='Inter, Segoe UI, sans-serif';

function unique(field){return [...new Set(DATA.map(r=>r[field]||'-'))].sort((a,b)=>String(a).localeCompare(String(b),'pt-BR'));}
function fillSelect(id,field){$(id).innerHTML='<option value="">Todos</option>'+unique(field).map(v=>`<option>${esc(v)}</option>`).join('');}
function configureDataControls(){fillSelect('tipoFilter','Tipo');fillSelect('situacaoFilter','Situacao');fillSelect('statusFilter','Status PCF');fillSelect('disciplinaFilter','Disciplina');fillSelect('responsavelFilter','Responsavel');const dates=DATA.map(r=>r['Data Recebimento']).filter(Boolean).sort();['dateFrom','dateTo'].forEach(id=>{$(id).value='';$(id).min=dates[0]||'';$(id).max=dates.at(-1)||'';});$('dataBase').textContent=dates.length?`Data-base: ${new Date(dates.at(-1)+'T12:00:00').toLocaleDateString('pt-BR')}`:'Data-base indisponível';}
configureDataControls();

function getFiltered(){const q=$('searchFilter').value.trim().toLocaleLowerCase('pt-BR');return DATA.filter(r=>(!$('tipoFilter').value||r.Tipo===$('tipoFilter').value)&&(!$('situacaoFilter').value||r.Situacao===$('situacaoFilter').value)&&(!$('statusFilter').value||r['Status PCF']===$('statusFilter').value)&&(!$('disciplinaFilter').value||r.Disciplina===$('disciplinaFilter').value)&&(!$('responsavelFilter').value||r.Responsavel===$('responsavelFilter').value)&&(!$('dateFrom').value||r['Data Recebimento']>=$('dateFrom').value)&&(!$('dateTo').value||r['Data Recebimento']<=$('dateTo').value)&&(!q||[r.Documento,r['Titulo do Documento'],r['PCF Recebida'],r['GRD Emissao']].some(v=>String(v||'').toLocaleLowerCase('pt-BR').includes(q))));}
function countBy(rows,field){return rows.reduce((a,r)=>{const k=r[field]||'-';a[k]=(a[k]||0)+1;return a;},{});}
function chart(id,type,data,options){if(charts[id])charts[id].destroy();charts[id]=new Chart($(id),{type,data,options:{responsive:true,maintainAspectRatio:false,interaction:{mode:'index',intersect:false},plugins:{legend:{labels:{usePointStyle:true,boxWidth:8}},tooltip:{backgroundColor:'#071521',borderColor:'#29485c',borderWidth:1,padding:12}},...options}});}
function updateKPIs(rows){const total=rows.length,overdue=rows.filter(r=>r.Situacao==='Vencida').length,open=rows.reduce((s,r)=>s+r.Open,0),answered=rows.filter(r=>r.Situacao==='Respondida'||r['PCF Respondida']).length,under=rows.reduce((s,r)=>s+r['Under Review'],0),critical=rows.filter(r=>r['Dias de Atraso']>90).length,avg=total?rows.reduce((s,r)=>s+Math.max(0,r['Dias de Atraso']),0)/total:0;$('kpiTotal').textContent=fmt.format(total);$('kpiTotalSub').textContent=`${fmt.format(critical)} com atraso superior a 90 d.u.`;$('kpiOverdue').textContent=fmt.format(overdue);$('kpiOverdueSub').textContent=`${pct(total?overdue/total*100:0)} da carteira · média ${avg.toFixed(1).replace('.',',')} d.u.`;$('kpiOpen').textContent=fmt.format(open);$('kpiOpenSub').textContent=`${fmt.format(under)} comentários em UNDER REVIEW`;$('kpiResponse').textContent=pct(total?answered/total*100:0);$('kpiResponseSub').textContent=`${fmt.format(answered)} respondidas de ${fmt.format(total)}`;$('activeCount').textContent=`${fmt.format(total)} de ${fmt.format(DATA.length)} PCFs no recorte`;$('executiveInsight').innerHTML=`<p class="text-xs font-black uppercase tracking-wider text-amber">Leitura executiva do recorte</p><p class="mt-2 text-lg font-bold">${total?`${pct(overdue/total*100)} da carteira está vencida, com ${fmt.format(open)} comentários OPEN e ${fmt.format(critical)} PCFs acima de 90 dias úteis.`:'Nenhum registro corresponde aos filtros selecionados.'}</p><p class="mt-2 text-sm text-slate-400">Prioridade sugerida: atacar documentos simultaneamente vencidos e com comentários OPEN, começando pelas disciplinas de maior concentração.</p>`;}
function updateCharts(rows){const situations=countBy(rows,'Situacao');chart('situationChart','doughnut',{labels:Object.keys(situations),datasets:[{data:Object.values(situations),backgroundColor:Object.keys(situations).map(k=>k==='Vencida'?colors.danger:k==='Respondida'?colors.mint:colors.amber),borderColor:'#0d2233',borderWidth:4,hoverOffset:8}]},{cutout:'66%',plugins:{legend:{position:'bottom'}}});
 const buckets=[['No prazo',r=>r['Dias de Atraso']<=0],['1–15',r=>r['Dias de Atraso']>=1&&r['Dias de Atraso']<=15],['16–30',r=>r['Dias de Atraso']>=16&&r['Dias de Atraso']<=30],['31–60',r=>r['Dias de Atraso']>=31&&r['Dias de Atraso']<=60],['61–90',r=>r['Dias de Atraso']>=61&&r['Dias de Atraso']<=90],['> 90',r=>r['Dias de Atraso']>90]];chart('agingChart','bar',{labels:buckets.map(x=>x[0]),datasets:[{label:'PCFs',data:buckets.map(x=>rows.filter(x[1]).length),backgroundColor:buckets.map((_,i)=>i===5?colors.danger:i>2?colors.amber:colors.cyan),borderRadius:7}]},{plugins:{legend:{display:false}},scales:{y:{beginAtZero:true,ticks:{precision:0}},x:{grid:{display:false}}}});
 const disciplines={};rows.forEach(r=>disciplines[r.Disciplina]=(disciplines[r.Disciplina]||0)+r.Open);const d=Object.entries(disciplines).sort((a,b)=>b[1]-a[1]).slice(0,8).reverse();chart('disciplineChart','bar',{labels:d.map(x=>x[0]),datasets:[{label:'Comentários OPEN',data:d.map(x=>x[1]),backgroundColor:colors.amber,borderRadius:6}]},{indexAxis:'y',plugins:{legend:{display:false}},scales:{x:{beginAtZero:true,ticks:{precision:0}},y:{grid:{display:false},ticks:{font:{size:10}}}}});
 const months={};rows.forEach(r=>{const m=(r['Data Recebimento']||'').slice(0,7);if(!m)return;months[m]??={total:0,vencidas:0};months[m].total++;if(r.Situacao==='Vencida')months[m].vencidas++;});const m=Object.keys(months).sort();chart('timelineChart','line',{labels:m.map(v=>new Date(v+'-02T12:00:00').toLocaleDateString('pt-BR',{month:'short',year:'2-digit'})),datasets:[{type:'bar',label:'Recebidas',data:m.map(v=>months[v].total),backgroundColor:'rgba(57,198,244,.42)',borderColor:colors.cyan,borderWidth:1,borderRadius:5},{type:'line',label:'Atualmente vencidas',data:m.map(v=>months[v].vencidas),borderColor:colors.danger,backgroundColor:colors.danger,tension:.32,pointRadius:3,pointHoverRadius:6}]},{scales:{y:{beginAtZero:true,ticks:{precision:0}},x:{grid:{display:false}}}});}
function updateTable(rows){const critical=[...rows].sort((a,b)=>(b['Dias de Atraso']-a['Dias de Atraso'])||(b.Open-a.Open)).slice(0,15);$('tableCount').textContent=`Exibindo ${critical.length} de ${rows.length} registros`;$('criticalTable').innerHTML=critical.map(r=>`<tr class="transition hover:bg-cyan/5"><td class="px-3 py-3"><div class="font-bold text-slate-100">${esc(r.Documento)}</div><div class="max-w-[390px] truncate text-xs text-slate-500" title="${esc(r['Titulo do Documento'])}">${esc(r['Titulo do Documento']||'-')}</div></td><td class="px-3 py-3"><span class="rounded-md bg-cyan/10 px-2 py-1 text-xs font-bold text-cyan">${esc(r.Tipo)}</span></td><td class="px-3 py-3 text-xs text-slate-300">${esc(r.Disciplina)}</td><td class="px-3 py-3"><span class="rounded-md px-2 py-1 text-xs font-bold ${r.Situacao==='Vencida'?'bg-danger/10 text-danger':'bg-amber/10 text-amber'}">${esc(r.Situacao)}</span></td><td class="px-3 py-3 text-right font-black text-danger">${fmt.format(r['Dias de Atraso'])}</td><td class="px-3 py-3 text-right font-black text-amber">${fmt.format(r.Open)}</td><td class="px-3 py-3 text-xs text-slate-300">${esc(r['Status PCF'])}</td><td class="px-3 py-3 text-xs text-slate-300">${esc(r.Responsavel)}</td></tr>`).join('')||'<tr><td colspan="8" class="px-3 py-10 text-center text-slate-500">Nenhum registro no recorte.</td></tr>';}
function update(){filtered=getFiltered();updateKPIs(filtered);updateCharts(filtered);updateTable(filtered);}
['tipoFilter','situacaoFilter','statusFilter','disciplinaFilter','responsavelFilter','dateFrom','dateTo'].forEach(id=>$(id).addEventListener('change',update));$('searchFilter').addEventListener('input',update);$('clearFilters').addEventListener('click',()=>{['tipoFilter','situacaoFilter','statusFilter','disciplinaFilter','responsavelFilter','dateFrom','dateTo','searchFilter'].forEach(id=>$(id).value='');update();});
function normalizeStatus(v){v=String(v||'-').trim().toUpperCase();return ['NOT RELESED','NOT RELEASED'].includes(v)?'NOT RELEASED':v;}
function excelDate(v){if(!v)return null;if(v instanceof Date&&!isNaN(v))return v.toISOString().slice(0,10);if(typeof v==='number'){const d=XLSX.SSF.parse_date_code(v);return d?`${d.y}-${String(d.m).padStart(2,'0')}-${String(d.d).padStart(2,'0')}`:null;}const s=String(v).trim();const br=s.match(/^(\d{1,2})\/(\d{1,2})\/(\d{4})/);if(br)return `${br[3]}-${br[2].padStart(2,'0')}-${br[1].padStart(2,'0')}`;return /^\d{4}-\d{2}-\d{2}/.test(s)?s.slice(0,10):null;}
function normalizeImported(row){const out={...row};['Data Recebimento','Data Resposta','Prazo (15 DU)'].forEach(k=>out[k]=excelDate(out[k]));['Dias sem Resposta','Dias de Atraso','Comentarios','Open','Under Review','Closed Calculado'].forEach(k=>out[k]=Number(out[k]||0));out['Status PCF']=normalizeStatus(out['Status PCF']);return out;}
$('loadExcel').addEventListener('click',()=>$('excelUpload').click());
$('excelUpload').addEventListener('change',async e=>{const file=e.target.files[0];if(!file)return;try{const wb=XLSX.read(await file.arrayBuffer(),{type:'array',cellDates:true});const name=wb.SheetNames.find(n=>n.trim().toLocaleLowerCase('pt-BR')==='dados pcfs');if(!name)throw new Error('A aba "Dados PCFs" não foi encontrada.');const rows=XLSX.utils.sheet_to_json(wb.Sheets[name],{defval:null,raw:true}).filter(r=>r.Documento).map(normalizeImported);if(!rows.length)throw new Error('A aba "Dados PCFs" não contém registros.');DATA=rows;$('sourceName').textContent=file.name;['tipoFilter','situacaoFilter','statusFilter','disciplinaFilter','responsavelFilter','dateFrom','dateTo','searchFilter'].forEach(id=>$(id).value='');configureDataControls();update();}catch(err){alert(`Não foi possível carregar o Excel. ${err.message}`);}finally{e.target.value='';}});
$('exportExcel').addEventListener('click',()=>{const cols=['Documento','Titulo do Documento','Tipo','Revisao','PCF Recebida','Data Recebimento','Resposta Esperada','PCF Respondida','Data Resposta','Situacao','Dias sem Resposta','Prazo (15 DU)','Dias de Atraso','Status PCF','Comentarios','Open','Under Review','Closed Calculado','Responsavel','GRD Emissao','Projeto','Disciplina','Link PCF','Link Documento'];const rows=filtered.map(r=>Object.fromEntries(cols.map(c=>[c,r[c]??''])));const ws=XLSX.utils.json_to_sheet(rows,{header:cols});ws['!autofilter']={ref:`A1:X${rows.length+1}`};ws['!freeze']={xSplit:0,ySplit:1};ws['!cols']=cols.map(c=>({wch:Math.min(45,Math.max(12,c.length+2))}));const wb=XLSX.utils.book_new();XLSX.utils.book_append_sheet(wb,ws,'Base Filtrada');const resumo=XLSX.utils.aoa_to_sheet([['Indicador','Valor'],['PCFs no recorte',filtered.length],['Vencidas',filtered.filter(r=>r.Situacao==='Vencida').length],['Comentários OPEN',filtered.reduce((s,r)=>s+r.Open,0)],['Respondidas',filtered.filter(r=>r.Situacao==='Respondida'||r['PCF Respondida']).length]]);XLSX.utils.book_append_sheet(wb,resumo,'Resumo');XLSX.writeFile(wb,'controle_respostas_pcf_recorte.xlsx');});
update();
</script>
</body></html>'''.replace("__DATA__", data_json)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(html, encoding="utf-8")
print(f"{OUTPUT}\nrecords={len(records)} bytes={OUTPUT.stat().st_size}")
