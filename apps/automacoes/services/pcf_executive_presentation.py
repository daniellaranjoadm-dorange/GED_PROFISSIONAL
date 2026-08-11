from io import BytesIO

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(7, 24, 39)
SURFACE = RGBColor(12, 36, 56)
CYAN = RGBColor(56, 189, 248)
BLUE = RGBColor(14, 165, 233)
WHITE = RGBColor(248, 251, 255)
TEXT = RGBColor(217, 231, 245)
MUTED = RGBColor(143, 168, 194)
RED = RGBColor(239, 83, 80)
AMBER = RGBColor(245, 158, 11)
GREEN = RGBColor(34, 197, 94)
LINE = RGBColor(39, 65, 89)


def _text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = str(value)
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def _rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    return shape


def _background(slide, presentation):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(.08))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = BLUE
    stripe.line.fill.background()


def _footer(slide, page, report_date):
    _text(slide, f"D'OR@NGE GED ENTERPRISE  |  {report_date:%d/%m/%Y}", .75, 7.08, 5.0, .18, 9, MUTED)
    _text(slide, f"{page:02d}", 12.25, 7.06, .45, .2, 9, CYAN, True, PP_ALIGN.RIGHT)


def _base_slide(presentation, headline, kicker, page, report_date):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _background(slide, presentation)
    _text(slide, kicker, .75, .32, 7.5, .25, 12, CYAN, True)
    _text(slide, headline, .75, .65, 11.8, .9, 30, WHITE, True)
    _footer(slide, page, report_date)
    return slide


def _chart_style(chart):
    chart.has_legend = False
    chart.has_title = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = LINE
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.font.color.rgb = TEXT
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.color.rgb = MUTED


def build_pcf_executive_presentation(records, summary, report_date):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1 — abertura
    slide = presentation.slides.add_slide(blank)
    _background(slide, presentation)
    _text(slide, "CONTROLE EXECUTIVO NAVAL", .8, .7, 6.0, .3, 14, CYAN, True)
    _text(slide, "Respostas PCF", .8, 1.45, 7.5, .85, 42, WHITE, True)
    _text(slide, "LD PROJETO BÁSICO", .8, 2.65, 5.5, .4, 20, CYAN, True)
    _text(slide, "Exposição contratual, comentários e prioridades de decisão", .8, 3.22, 9.0, .55, 22, MUTED)
    _text(slide, "SLA: 15 dias úteis após o recebimento", .8, 5.1, 5.2, .4, 19, AMBER, True)
    _text(slide, f"{summary['total']} PCFs no escopo  |  Data-base: {report_date:%d/%m/%Y}", .8, 5.58, 6.2, .3, 15, TEXT)
    _text(slide, "DIRETORIA", 10.7, 5.9, 1.7, .3, 13, CYAN, True, PP_ALIGN.RIGHT)
    _footer(slide, 1, report_date)

    # 2 — síntese executiva
    slide = _base_slide(presentation, "A exposição exige ação imediata", "SÍNTESE EXECUTIVA", 2, report_date)
    metrics = [
        ("VENCIDAS", summary["vencidas"], RED),
        ("AGUARDANDO", summary["aguardando"], AMBER),
        ("RESPONDIDAS", summary["respondidas"], GREEN),
        ("COMENTÁRIOS", summary["comentarios_total"], CYAN),
    ]
    for index, (label, value, color) in enumerate(metrics):
        x = .75 + index * 3.0
        _text(slide, label, x, 1.72, 2.5, .3, 12, color, True)
        _text(slide, value, x, 2.08, 2.5, .62, 30, WHITE, True)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.78), Inches(2.38), Inches(.04))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    _text(slide, f"{summary['exposicao_pct']}% da carteira está vencida.", .75, 3.55, 7.0, .5, 27, RED, True)
    _text(slide, f"Comentários: {summary['comentarios_abertos']} OPEN  |  {summary['comentarios_revisao']} UNDER REVIEW  |  {summary['comentarios_fechados']} CLOSED calculado", .75, 4.12, 11.6, .36, 17, TEXT, True)
    _text(slide, "O aging utiliza exclusivamente a Data Recebimento da PCF.", .75, 4.52, 10.0, .3, 15, MUTED)
    _rect(slide, .7, 5.18, 11.95, .92, SURFACE, LINE)
    _text(slide, "DECISÃO", 1.02, 5.38, 1.35, .3, 12, CYAN, True)
    _text(slide, "Priorizar PCFs vencidas com comentários OPEN e compromisso nominal de resposta.", 2.7, 5.3, 9.4, .44, 17, WHITE, True)

    # 3 — concentração técnica
    discipline = {}
    for item in records:
        discipline[item["disciplina"]] = discipline.get(item["disciplina"], 0) + 1
    discipline_items = sorted(discipline.items(), key=lambda pair: (-pair[1], pair[0]))[:8]
    slide = _base_slide(presentation, "A exposição está concentrada em poucas disciplinas críticas", "CONCENTRAÇÃO TÉCNICA", 3, report_date)
    data = ChartData()
    data.categories = [label for label, _ in discipline_items]
    data.add_series("PCFs", [value for _, value in discipline_items])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.65), Inches(1.72), Inches(8.15), Inches(4.65), data).chart
    _chart_style(chart)
    chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = BLUE
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.show_value = True
    chart.plots[0].data_labels.font.color.rgb = WHITE
    chart.plots[0].data_labels.font.size = Pt(11)
    chart.plots[0].data_labels.font.bold = True
    top_three = discipline_items[:3]
    top_total = sum(value for _, value in top_three)
    _text(slide, "FOCO DE GESTÃO", 9.25, 1.8, 2.6, .3, 13, CYAN, True)
    _text(slide, f"{top_total} PCFs", 9.25, 2.32, 2.6, .55, 32, WHITE, True)
    _text(slide, "nas três disciplinas com maior concentração", 9.25, 2.98, 2.8, .65, 16, TEXT)
    _text(slide, "\n\n".join(f"{value}  {label}" for label, value in top_three), 9.25, 3.8, 3.2, 1.65, 16, WHITE, True)
    _text(slide, "Direcionar cobrança e capacidade técnica por disciplina reduz dispersão.", 9.25, 5.7, 3.15, .62, 14, MUTED)

    # 4 — aging e comentários
    aging_order = [("No prazo", 0, 0), ("1–15", 1, 15), ("16–30", 16, 30), ("31–60", 31, 60), ("61–90", 61, 90), (">90", 91, None)]
    aging_values = []
    for _, minimum, maximum in aging_order:
        count = sum(1 for item in records if item["dias_atraso"] >= minimum and (maximum is None or item["dias_atraso"] <= maximum))
        aging_values.append(count)
    slide = _base_slide(presentation, f"Há {aging_values[-1]} PCFs com atraso superior a 90 dias úteis", "AGING E COMENTÁRIOS", 4, report_date)
    data = ChartData(); data.categories = [item[0] for item in aging_order]; data.add_series("PCFs", aging_values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.72), Inches(7.6), Inches(4.6), data).chart
    _chart_style(chart); chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = RED
    chart.plots[0].has_data_labels = True; chart.plots[0].data_labels.show_value = True
    chart.plots[0].data_labels.font.color.rgb = WHITE
    chart.plots[0].data_labels.font.size = Pt(11)
    chart.plots[0].data_labels.font.bold = True
    _text(slide, "Comentários das PCFs", 8.9, 1.75, 3.2, .35, 18, CYAN, True)
    comment_rows = [("OPEN", summary["comentarios_abertos"], AMBER), ("UNDER REVIEW", summary["comentarios_revisao"], CYAN), ("CLOSED CALCULADO", summary["comentarios_fechados"], GREEN)]
    for index, (label, value, color) in enumerate(comment_rows):
        y = 2.45 + index * 1.25
        _text(slide, label, 9.15, y, 2.2, .3, 14, color, True)
        _text(slide, value, 11.0, y-.16, 1.2, .6, 28, WHITE, True, PP_ALIGN.RIGHT)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.15), Inches(y+.62), Inches(3.0), Inches(.04))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    if summary["comentarios_inconsistentes"]:
        quantidade = summary["comentarios_inconsistentes"]
        rotulo = "inconsistência" if quantidade == 1 else "inconsistências"
        _text(slide, f"{quantidade} {rotulo} na fonte sinalizada", 9.15, 5.98, 3.0, .42, 11, RED, True)

    # 5 — prioridades
    slide = _base_slide(presentation, "Oito documentos concentram a maior criticidade", "PRIORIDADES DE COBRANÇA", 5, report_date)
    top = sorted(records, key=lambda item: (item["dias_atraso"], item["open_comments"]), reverse=True)[:8]
    columns = [(.75, 3.45), (4.25, 3.75), (8.25, 1.25), (9.72, 2.7)]
    for (label, (x, width)) in zip(("DOCUMENTO / TÍTULO", "PCF RECEBIDA", "ATRASO", "COMENTÁRIOS — OPEN / UNDER REVIEW / CLOSED"), columns):
        _text(slide, label, x, 1.56, width, .4, 10 if x < 9 else 9, CYAN, True)
    for index, item in enumerate(top):
        y = 2.05 + index * .52
        if index % 2 == 0:
            _rect(slide, .62, y-.05, 12.05, .48, SURFACE)
        title = str(item["titulo"] or "-")
        if len(title) > 36:
            title = title[:35] + "…"
        _text(slide, f"{item['documento']}\n{title}", .75, y, 3.55, .4, 10, WHITE, True)
        _text(slide, item["pcf"], 4.25, y+.06, 3.7, .27, 9, TEXT)
        _text(slide, f"{item['dias_atraso']} d.u.", 8.25, y+.06, 1.2, .27, 12, RED, True)
        _text(slide, f"{item['open_comments']} / {item['under_review']} / {item['closed_comments']}", 9.72, y+.06, 2.65, .27, 12, AMBER, True)

    # 6 — composição da carteira
    document_types = {}
    disciplines = {}
    for item in records:
        document_type = item.get("tipo_documento") or "NÃO INFORMADO"
        discipline = item.get("disciplina") or "NÃO INFORMADA"
        document_types[document_type] = document_types.get(document_type, 0) + 1
        disciplines[discipline] = disciplines.get(discipline, 0) + 1
    type_items = sorted(document_types.items(), key=lambda pair: (-pair[1], pair[0]))[:6] or [("SEM DADOS", 0)]
    discipline_items = sorted(disciplines.items(), key=lambda pair: (-pair[1], pair[0]))[:6] or [("SEM DADOS", 0)]

    slide = _base_slide(presentation, "Tipos documentais e disciplinas definem o foco da carteira", "COMPOSIÇÃO DA CARTEIRA", 6, report_date)
    _text(slide, "TIPOS DE DOCUMENTO", .72, 1.5, 5.7, .3, 13, CYAN, True)
    _text(slide, "DISCIPLINAS", 6.85, 1.5, 5.7, .3, 13, CYAN, True)

    for items, x, color in ((type_items, .65, BLUE), (discipline_items, 6.78, CYAN)):
        ranked = list(reversed(items))
        data = ChartData()
        data.categories = [label for label, _ in ranked]
        data.add_series("PCFs", [value for _, value in ranked])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(1.88), Inches(5.85), Inches(4.72), data).chart
        _chart_style(chart)
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = color
        chart.plots[0].has_data_labels = True
        chart.plots[0].data_labels.show_value = True
        chart.plots[0].data_labels.font.color.rgb = WHITE
        chart.plots[0].data_labels.font.size = Pt(11)
        chart.plots[0].data_labels.font.bold = True

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
