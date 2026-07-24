"""Geração segura do Relatório Executivo KM.

A fonte é aberta somente para leitura pelo extrator. Esta rotina nunca grava na LD:
ela cria um novo XLSX versionado e um log de auditoria a cada execução.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import traceback
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PASTA_BASE = Path(
    r"\\virm-rgr022\FILESERVER\Projetos\05_HANDYMAX\09. Doc Control\3 - LD"
)
PASTA_SAIDA = PASTA_BASE / "Relatorios Executivos"
PASTA_LOGS = PASTA_BASE / "Logs"
EXTRATOR = Path(__file__).with_name("relatorio_executivo_km_dados.py")

AZUL = "17365D"
AZUL_MEDIO = "2F75B5"
AZUL_CLARO = "D9EAF7"
VERDE = "70AD47"
AMARELO = "FFC000"
VERMELHO = "C00000"
CINZA = "E7E6E6"
BRANCO = "FFFFFF"
BORDA = Side(style="thin", color="B7B7B7")


def _texto(valor):
    return "" if valor is None else str(valor)


def _titulo(ws, texto, subtitulo=""):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:L1")
    ws["A1"] = texto
    ws["A1"].font = Font(size=20, bold=True, color=BRANCO)
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    if subtitulo:
        ws.merge_cells("A2:L2")
        ws["A2"] = subtitulo
        ws["A2"].font = Font(size=10, italic=True, color="555555")
        ws.row_dimensions[2].height = 22


def _kpi(ws, celula, rotulo, valor, cor=AZUL_MEDIO):
    c = ws[celula]
    c.value = f"{rotulo}\n{valor}"
    c.font = Font(size=13, bold=True, color=BRANCO)
    c.fill = PatternFill("solid", fgColor=cor)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
    ws.row_dimensions[c.row].height = 48


def _valor(item, chave):
    valor = item.get(chave, "")
    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False)
    return valor


def _adicionar_tabela(ws, inicio, colunas, registros, nome):
    linha, coluna = inicio
    for deslocamento, (chave, cabecalho) in enumerate(colunas):
        c = ws.cell(linha, coluna + deslocamento, cabecalho)
        c.font = Font(bold=True, color=BRANCO)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
    for i, item in enumerate(registros, start=linha + 1):
        for deslocamento, (chave, _cabecalho) in enumerate(colunas):
            c = ws.cell(i, coluna + deslocamento, _valor(item, chave))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
            if chave == "prioridade":
                texto = _texto(c.value).upper()
                cor = VERMELHO if "CR" in texto else AMARELO if "ALTA" in texto else CINZA
                c.fill = PatternFill("solid", fgColor=cor)
                if cor == VERMELHO:
                    c.font = Font(color=BRANCO, bold=True)
    fim = max(linha + 1, linha + len(registros))
    # Tabelas do Excel precisam de ao menos uma linha de dados.
    if not registros:
        ws.cell(linha + 1, coluna, "Sem registros para os critérios atuais")
    ref = f"{get_column_letter(coluna)}{linha}:{get_column_letter(coluna + len(colunas) - 1)}{fim}"
    tabela = Table(displayName=nome, ref=ref)
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    ws.add_table(tabela)
    ws.freeze_panes = ws.cell(linha + 1, coluna)
    # A própria Tabela do Excel já possui AutoFilter. Não criar um segundo
    # filtro de planilha sobre a mesma faixa: o Excel considera a sobreposição
    # inválida e remove as tabelas ao tentar reparar o arquivo.
    for deslocamento, (chave, cabecalho) in enumerate(colunas):
        largura = min(48, max(12, len(cabecalho) + 2))
        if chave in {"titulo", "titulos", "kms", "dox", "action", "correlacao"}:
            largura = 36
        ws.column_dimensions[get_column_letter(coluna + deslocamento)].width = largura


def _aba_dados(wb, nome, titulo, registros, colunas, tabela):
    ws = wb.create_sheet(nome)
    _titulo(ws, titulo, "Fonte: LD PROJETO BASICO e GENERAL LIST KM | geração somente leitura")
    _adicionar_tabela(ws, (4, 1), colunas, registros, tabela)
    return ws


def _criar_workbook(payload):
    r = payload["resumo"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo Executivo"
    _titulo(ws, "RELATÓRIO EXECUTIVO — PROJETO BÁSICO / KM", f"Data-base: {r['data_referencia']} | Arquivo fonte: {r['arquivo_fonte']}")

    kpis = [
        ("A4", "TP únicos pendentes", r["tp_unicos_pendentes"], VERMELHO),
        ("C4", "Docs KM correlacionados", r["correlacionados_pendentes"], AZUL_MEDIO),
        ("E4", "ETS — escopo LD", r["ld_ets_linhas"], AZUL_MEDIO),
        ("G4", "Armature — escopo LD", r["ld_armature_linhas"], AZUL_MEDIO),
        ("I4", "PCFs abertas", r["pcf_abertas_unicas"], AMARELO),
        ("K4", "Pendências BV", r["bv_pendencias_total"], VERMELHO),
    ]
    for celula, rotulo, valor, cor in kpis:
        _kpi(ws, celula, rotulo, valor, cor)
        ws.merge_cells(start_row=4, start_column=ws[celula].column, end_row=4, end_column=ws[celula].column + 1)

    ws["A7"] = "Leitura executiva"
    ws["A7"].font = Font(size=14, bold=True, color=AZUL)
    notas = [
        f"• {r['tp_unicos_pendentes']} números Transpetro únicos aguardam emissão, ligados a {r['correlacionados_pendentes']} documentos KM.",
        f"• ETS: {r['ld_ets_recebidas_pendentes']} recebidas e não emitidas; {r['ld_ets_nao_recebidas']} ainda não recebidas.",
        f"• Armature List: {r['ld_armature_recebidas_pendentes']} recebidas e não emitidas; {r['ld_armature_nao_recebidas']} ainda não recebidas.",
        f"• {r['general_tbd_unicos']} documentos GENERAL LIST estão como TBD ({r['general_tbd_recebidos']} recebidos).",
        f"• {r['pcf_abertas_unicas']} PCFs sem resposta; {r['pcf_criticas']} críticas e {r['pcf_com_comentarios_abertos']} com comentários abertos.",
        f"• {r['bv_documentos_unicos']} documentos foram enviados ao BV; {r['bv_com_pendencias']} possuem pendências.",
        f"• {r['ld_kongsberg_sem_general']} números TP no escopo Kongsberg não possuem correlação na GENERAL LIST KM.",
    ]
    for i, nota in enumerate(notas, start=8):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
        ws.cell(i, 1, nota)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[i].height = 24

    ws["A17"] = "Reconciliação dos universos"
    ws["A17"].font = Font(size=14, bold=True, color=AZUL)
    reconciliacao = [
        {"universo": "ETS", "escopo_ld": r["ld_ets_linhas"], "recebido_pendente": r["ld_ets_recebidas_pendentes"], "nao_recebido": r["ld_ets_nao_recebidas"], "emitido": r["ld_ets_emitidas"]},
        {"universo": "Armature List", "escopo_ld": r["ld_armature_linhas"], "recebido_pendente": r["ld_armature_recebidas_pendentes"], "nao_recebido": r["ld_armature_nao_recebidas"], "emitido": r["ld_armature_emitidas"]},
    ]
    _adicionar_tabela(ws, (18, 1), [("universo", "Universo"), ("escopo_ld", "Escopo LD"), ("recebido_pendente", "Recebido e não emitido"), ("nao_recebido", "Não recebido"), ("emitido", "Emitido")], reconciliacao, "tbReconciliacao")

    ws["G17"] = "Controles de qualidade"
    ws["G17"].font = Font(size=14, bold=True, color=AZUL)
    controles = [
        ("Recebidos únicos na GENERAL LIST", r["recebidos_unicos"]),
        ("Duplicidades consolidadas", r["duplicidades_consolidadas"]),
        ("TP distintos válidos na LD", r["tp_distintos_resumo"]),
        ("KM sem correlação TP", r["sem_correlacao_tp"]),
    ]
    for i, (rotulo, valor) in enumerate(controles, start=18):
        ws.cell(i, 7, rotulo)
        ws.cell(i, 10, valor)
        ws.merge_cells(start_row=i, start_column=7, end_row=i, end_column=9)
        for c in range(7, 11):
            ws.cell(i, c).border = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)
        ws.cell(i, 10).font = Font(bold=True, color=AZUL)

    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.freeze_panes = "A4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    comum = [("km", "Nº Documento KM"), ("titulo", "Título"), ("disciplina", "Disciplina"), ("customer", "Nº Transpetro/Customer"), ("transmittal", "Transmittal"), ("data", "Data recebimento"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("status_emissao", "Status emissão"), ("linhas_ld", "Linhas LD")]
    _aba_dados(wb, "Pend Corr TP", "PENDÊNCIAS CORRELACIONADAS — AGRUPADAS POR Nº TRANSPETRO", payload["correlacionados_tp_agrupados"], [("tp", "Nº Transpetro"), ("qtd_km", "Qtd. Docs KM"), ("kms", "Documentos KM"), ("titulos", "Títulos"), ("disciplinas", "Disciplinas"), ("transmittals", "Transmittals"), ("primeira_data", "Primeiro recebimento"), ("ultima_data", "Último recebimento"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("status_emissao", "Status emissão"), ("linhas_ld", "Linhas LD")], "tbPendCorrTP")
    _aba_dados(wb, "ETS Escopo", "ETS — UNIVERSO OFICIAL DA LD", payload["escopo_ets"], [("km", "Nº KM"), ("titulo", "Título"), ("dox", "Nº interno DOX"), ("disciplina", "Disciplina"), ("transmittal", "Transmittal"), ("data", "Data recebimento"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("status", "Status"), ("status_emissao", "Status emissão"), ("situacao", "Situação"), ("linha_ld", "Linha LD")], "tbETS")
    _aba_dados(wb, "Armature List", "ARMATURE LIST — UNIVERSO OFICIAL DA LD", payload["escopo_armature"], [("km", "Nº KM"), ("titulo", "Título"), ("dox", "Nº interno DOX"), ("disciplina", "Disciplina"), ("transmittal", "Transmittal"), ("data", "Data recebimento"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("status", "Status"), ("status_emissao", "Status emissão"), ("situacao", "Situação"), ("linha_ld", "Linha LD")], "tbArmature")
    _aba_dados(wb, "KM Sem Correlação", "DOCUMENTOS KM SEM CORRELAÇÃO TRANSPETRO", payload["sem_correlacao"], comum, "tbSemCorrelacao")
    _aba_dados(wb, "GENERAL TBD", "GENERAL LIST KM — CUSTOMER TBD", payload["tbd_detalhe"], [("km", "Nº KM"), ("titulo", "Título"), ("disciplina", "Disciplina"), ("toc", "Tipo"), ("transmittal", "Transmittal"), ("data", "Data"), ("recebido", "Recebido"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("linha_gl", "Linha GENERAL")], "tbTBD")
    _aba_dados(wb, "PCFs em Aberto", "PCFs EMITIDAS SEM PCF RESPONDIDA", payload["pcf_abertas"], [("pcf", "PCF Nº"), ("tp", "Nº Transpetro"), ("dox", "Nº DOX"), ("revisao", "Revisão"), ("titulo", "Título"), ("responsavel", "Responsável"), ("data", "Data PCF"), ("dias", "Dias em aberto"), ("prioridade", "Prioridade"), ("status_pcf", "Status PCF"), ("qtd_comentarios", "Comentários"), ("comentarios_abertos", "Abertos"), ("em_revisao", "Em revisão"), ("status_final", "Status final"), ("linhas_ld", "Linhas LD")], "tbPCFAbertas")
    _aba_dados(wb, "Kongsberg para BV", "DOCUMENTOS KONGSBERG EMITIDOS PARA O BV", payload["documentos_bv"], [("km", "Nº KM"), ("titulo", "Título"), ("disciplina", "Disciplina"), ("customer", "Nº Transpetro"), ("posted", "Postado BV"), ("status", "Status BV"), ("since", "Desde"), ("action", "Ação"), ("pending", "Pendências"), ("dias", "Dias"), ("prioridade", "Prioridade"), ("linha_gl", "Linha GENERAL")], "tbBV")
    _aba_dados(wb, "Resumo Status", "STATUS DA LD — CONTAGEM DISTINTA DE Nº TRANSPETRO", payload["matriz_status"], [("responsavel", "Responsável pela emissão"), ("status", "Status"), ("qtd_tp", "Qtd. TP distintos")], "tbStatus")
    _aba_dados(wb, "LD Kongsberg sem GL", "ESCOPO KONGSBERG NA LD SEM CORRELAÇÃO NA GENERAL LIST", payload["ld_kongsberg_sem_general"], [("tp", "Nº Transpetro"), ("dox", "Nº DOX"), ("revisao", "Revisão"), ("titulo", "Título"), ("disciplina", "Disciplina"), ("status", "Status"), ("status_emissao", "Status emissão"), ("resp_issue", "Responsável"), ("linha_ld", "Linha LD")], "tbSemGL")

    regras = wb.create_sheet("Regras e Auditoria")
    _titulo(regras, "REGRAS, RASTREABILIDADE E LIMITAÇÕES")
    linhas = [
        "1. A LD é aberta somente para leitura; o relatório não grava nem recalcula a planilha fonte.",
        "2. Documentos KM duplicados são consolidados pelo registro com Data de Recebimento mais recente e, em empate, maior Transmittal.",
        "3. A aba Pend Corr TP possui uma linha por número Transpetro único e pode agrupar vários documentos KM.",
        "4. ETS e Armature List usam o universo oficial da LD: coluna C = NOT APPLICABLE e tipos ET/LI.",
        "5. KM sem correlação não é sinônimo de atraso confirmado; exige saneamento da chave documental.",
        "6. PCF em aberto: PCF Nº preenchida e PCF RESPONDIDA vazia, consolidada por número de PCF.",
        "7. Registros Kongsberg sem GENERAL ignoram números Transpetro inválidos, inclusive NOT APPLICABLE.",
        f"8. Arquivo fonte: {r['arquivo_fonte']}",
        f"9. Data de referência: {r['data_referencia']}",
    ]
    for i, linha in enumerate(linhas, start=4):
        regras.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
        regras.cell(i, 1, linha)
        regras.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        regras.row_dimensions[i].height = 32
    regras.column_dimensions["A"].width = 26
    return wb


def executar():
    inicio = datetime.now()
    carimbo = inicio.strftime("%Y%m%d_%H%M%S")
    nome = f"RELATORIO_EXECUTIVO_KM_{carimbo}.xlsx"
    log = PASTA_LOGS / f"RELATORIO_EXECUTIVO_KM_{carimbo}.log"
    final = PASTA_SAIDA / nome
    dados_json = Path(tempfile.gettempdir()) / f"dados_relatorio_executivo_{carimbo}_{os.getpid()}.json"
    linhas_log = [f"Início: {inicio:%d/%m/%Y %H:%M:%S}", "Modo: SOMENTE LEITURA", f"Fonte: {PASTA_BASE}"]
    try:
        PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
        PASTA_LOGS.mkdir(parents=True, exist_ok=True)
        if not EXTRATOR.exists():
            raise FileNotFoundError(f"Extrator não encontrado: {EXTRATOR}")
        ambiente = os.environ.copy()
        ambiente["RELATORIO_EXECUTIVO_DADOS"] = str(dados_json)
        processo = subprocess.run([sys.executable, str(EXTRATOR)], capture_output=True, text=True, timeout=420, env=ambiente)
        if processo.returncode != 0:
            raise RuntimeError(f"Falha na leitura da LD: {processo.stderr[-4000:]}")
        payload = json.loads(dados_json.read_text(encoding="utf-8"))
        wb = _criar_workbook(payload)
        # O temporário fica no mesmo compartilhamento para permitir renomeação
        # atômica no Windows (os.replace não cruza unidades/volumes).
        temporario = PASTA_SAIDA / f".{nome}.partial"
        wb.save(temporario)
        wb.close()
        # Publicação atômica: arquivo incompleto nunca aparece na pasta final.
        os.replace(temporario, final)
        tamanho = final.stat().st_size
        if tamanho < 10_000:
            raise RuntimeError(f"Arquivo gerado com tamanho inesperado: {tamanho} bytes")
        fim = datetime.now()
        linhas_log += [f"Fim: {fim:%d/%m/%Y %H:%M:%S}", f"Arquivo: {final}", f"Tamanho: {tamanho} bytes", "Status: SUCESSO"]
        log.write_text("\n".join(linhas_log), encoding="utf-8")
        return {
            "ok": True,
            "mensagem": f"Relatório executivo gerado com segurança: {nome}",
            "quantidade_processada": payload["resumo"]["recebidos_unicos"],
            "detalhes": {"arquivo": str(final), "log": str(log), "modo": "somente leitura"},
        }
    except Exception as exc:
        linhas_log += [f"Status: FALHA", f"Erro: {exc}", traceback.format_exc()]
        try:
            PASTA_LOGS.mkdir(parents=True, exist_ok=True)
            log.write_text("\n".join(linhas_log), encoding="utf-8")
        except Exception:
            pass
        return {"ok": False, "mensagem": f"Falha ao gerar relatório executivo: {exc}", "quantidade_processada": 0, "detalhes": {"log": str(log)}}
    finally:
        try:
            dados_json.unlink(missing_ok=True)
        except Exception:
            pass
