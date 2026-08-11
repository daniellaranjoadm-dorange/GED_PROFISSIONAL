import os
import tempfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import Workbook

from apps.automacoes.services.atualizar_ld_projeto_basico import (
    _pcf_data_documental_arquivo,
    _pcf_qtd_e_open_comments_timeline,
    _pcf_status_final_timeline,
    _valor_km_informado,
    indexar_pcfs,
    status_final_da_pcf,
)


class PCFCommentsReaderTests(SimpleTestCase):
    def test_na_is_not_receipt_evidence_for_km(self):
        for value in (None, "", "-", "N/A", "NA", "#N/A", "NOT APPLICABLE"):
            self.assertFalse(_valor_km_informado(value))
        self.assertTrue(_valor_km_informado("T-44943"))
        self.assertTrue(_valor_km_informado(datetime(2026, 7, 26)))

    def test_status_reader_ignores_excel_errors_in_comment_area(self):
        ws = Workbook().active
        ws["A6"] = "PCF History"
        ws["E7"] = "Status"
        ws["E9"] = "NOT RELEASED"
        ws["A18"] = "Item"
        ws["E28"] = "#VALUE!"
        ws["E46"] = "#VALUE!"

        self.assertEqual(_pcf_status_final_timeline(ws), "NOT RELEASED")

    def test_timeline_fallback_rejects_invalid_status(self):
        key = "PCF-I-DE-4880.00-4034-610-CZ1-001_R0"
        self.assertEqual(status_final_da_pcf({key: "#VALUE!"}, key), "")
        self.assertEqual(
            status_final_da_pcf({key: "NOT RELESED"}, key),
            "NOT RELEASED",
        )

    def test_indexer_normalizes_document_date_and_file_datetime(self):
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            primeira = raiz / "a" / "PCF-I-PT-4880.00-0001-000-CZ1-001_R0.xlsx"
            segunda = raiz / "b" / "PCF-I-PT-4880.00-0001-000-CZ1-001_R0.xlsx"
            primeira.parent.mkdir()
            segunda.parent.mkdir()
            primeira.touch()
            segunda.touch()

            def data_documental(caminho):
                return date(2026, 8, 1) if Path(caminho).parent.name == "a" else None

            with patch(
                "apps.automacoes.services.atualizar_ld_projeto_basico._pcf_data_documental_arquivo",
                side_effect=data_documental,
            ), patch(
                "apps.automacoes.services.atualizar_ld_projeto_basico._file_datetime",
                return_value=datetime(2026, 8, 2, 17, 30),
            ):
                indice = indexar_pcfs(raiz)

        registro = indice["I-PT-4880.00-0001-000-CZ1-001"]["0"]
        self.assertEqual(registro["date"], date(2026, 8, 2))
        self.assertEqual(Path(registro["path"]).parent.name, "b")

    def test_uses_latest_open_comments_from_pcf_history(self):
        ws = Workbook().active
        ws["A1"] = "PCF History"
        ws["A3"] = "Round"
        ws["H3"] = "Open Comments"
        ws["A4"], ws["H4"] = "a)", 16
        ws["A5"], ws["H5"] = "b)", 19
        ws["A6"], ws["H6"] = "c)", 3

        ws["J12"] = "Comment Status"
        for row in range(13, 46):
            ws.cell(row, 10).value = "OPEN" if row < 29 else "CLOSED"

        total, open_comments, under_review = (
            _pcf_qtd_e_open_comments_timeline(ws)
        )

        self.assertEqual(total, 33)
        self.assertEqual(open_comments, 3)
        self.assertEqual(under_review, 0)

    def test_falls_back_to_detailed_status_count_without_history(self):
        ws = Workbook().active
        ws["C2"] = "Comment Status"
        ws["C3"] = "OPEN"
        ws["C4"] = "CLOSED"
        ws["C5"] = "OPEN"
        ws["C6"] = "UNDER REVIEW"

        total, open_comments, under_review = (
            _pcf_qtd_e_open_comments_timeline(ws)
        )

        self.assertEqual(total, 4)
        self.assertEqual(open_comments, 2)
        self.assertEqual(under_review, 1)

    def test_accepts_zero_as_latest_history_balance(self):
        ws = Workbook().active
        ws["A1"] = "PCF History"
        ws["A3"] = "Round"
        ws["H3"] = "Open Comments"
        ws["A4"], ws["H4"] = "a)", 5
        ws["A5"], ws["H5"] = "b)", 0
        ws["J12"] = "Comment Status"
        ws["J13"] = "OPEN"

        _, open_comments, _ = _pcf_qtd_e_open_comments_timeline(ws)

        self.assertEqual(open_comments, 0)

    def test_parses_composite_balance_from_latest_history_round(self):
        ws = Workbook().active
        ws["A1"] = "PCF History"
        ws["A3"] = "Round"
        ws["J3"] = "Open Comments"
        ws["A4"], ws["J4"] = "a)", 15
        ws["A5"], ws["J5"] = "b)", 7
        ws["A6"], ws["J6"] = "c)", "2 OPEN, 4 UNDER REVIEW"
        ws["L10"] = "Comment Status"
        for row, status in enumerate(
            ["OPEN"] * 7 + ["UNDER REVIEW"] * 4 + ["CLOSED"] * 4,
            start=11,
        ):
            ws.cell(row, 12).value = status

        total, open_comments, under_review = (
            _pcf_qtd_e_open_comments_timeline(ws)
        )

        self.assertEqual(total, 15)
        self.assertEqual(open_comments, 2)
        self.assertEqual(under_review, 4)

    def test_reads_document_date_from_pcf_header(self):
        wb = Workbook()
        ws = wb.active
        ws["K5"] = "Date"
        ws["L5"] = datetime(2026, 2, 25)

        handle, caminho = tempfile.mkstemp(suffix=".xlsx")
        os.close(handle)
        try:
            wb.save(caminho)
            self.assertEqual(
                _pcf_data_documental_arquivo(caminho),
                date(2026, 2, 25),
            )
        finally:
            wb.close()
            os.unlink(caminho)
