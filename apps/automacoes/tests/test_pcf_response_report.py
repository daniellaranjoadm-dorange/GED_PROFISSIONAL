from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase
from pptx import Presentation

from apps.automacoes.services.pcf_executive_presentation import (
    build_pcf_executive_presentation,
)
from apps.automacoes.services.pcf_response_report import build_record, summarize
from apps.automacoes import views


def _document(**overrides):
    values = {
        "pk": 1,
        "data_pcf": date(2026, 1, 5),
        "data_resposta": None,
        "pcf": "PCF-I-DE-4880.00-0001-000-CZ1-001_R0",
        "pcf_resposta": "",
        "revisao": "0",
        "open_comments": 6,
        "under_review": 0,
        "qtd_comentarios": 5,
        "origem_aba": "LD Projeto Basico",
        "casco": "",
        "disciplina": "MAQUINAS",
        "documento": "I-DE-4880.00-0001-000-CZ1-001",
        "titulo": "DOCUMENTO DE TESTE",
        "status_final_pcf": "NOT RELEASED",
        "status": "",
        "grd": "GRD-001",
        "caminho_pcf": "",
        "caminho_documento": "",
    }
    values.update(overrides)
    return SimpleNamespace(**values)


class PCFResponseReportTests(SimpleTestCase):
    def test_aging_uses_received_pcf_date_and_reconciles_comments(self):
        record = build_record(_document(), today=date(2026, 2, 10))

        self.assertEqual(record["data_recebimento"], date(2026, 1, 5))
        self.assertEqual(record["prazo"], date(2026, 1, 26))
        self.assertGreater(record["dias_atraso"], 0)
        self.assertEqual(record["responsavel"], "McLaren")
        self.assertEqual(record["closed_comments"], -1)
        self.assertFalse(record["comments_consistent"])

    def test_summary_reconciles_open_under_review_and_closed(self):
        records = [
            build_record(_document(), today=date(2026, 2, 10)),
            build_record(
                _document(
                    pk=2,
                    documento="I-PT-4880.00-0002-000-CZ1-001",
                    pcf="PCF-I-PT-4880.00-0002-000-CZ1-001_R0",
                    qtd_comentarios=10,
                    open_comments=2,
                    under_review=3,
                ),
                today=date(2026, 2, 10),
            ),
        ]
        result = summarize(records)

        self.assertEqual(result["comentarios_total"], 15)
        self.assertEqual(result["comentarios_abertos"], 8)
        self.assertEqual(result["comentarios_revisao"], 3)
        self.assertEqual(result["comentarios_fechados"], 4)
        self.assertEqual(result["comentarios_reconciliacao"], 0)
        self.assertEqual(result["comentarios_inconsistentes"], 1)

    def test_historical_response_is_hidden_and_does_not_close_latest_cycle(self):
        record = build_record(
            _document(
                documento="I-PT-4880.00-2493-000-CZ1-001",
                pcf="PCF-I-PT-4880.00-2493-000-CZ1-001_R0D",
                data_pcf=date(2026, 6, 12),
                pcf_resposta="PCF-I-PT-4880.00-2493-000-CZ1-001_R0C",
                data_resposta=date(2026, 6, 2),
            ),
            today=date(2026, 8, 3),
        )

        self.assertEqual(record["pcf"], "PCF-I-PT-4880.00-2493-000-CZ1-001_R0D")
        self.assertEqual(record["pcf_resposta"], "")
        self.assertIsNone(record["data_resposta"])
        self.assertEqual(record["resposta_esperada"], "R0E")
        self.assertFalse(record["resposta_ciclo_atual"])
        self.assertEqual(record["situacao"], "Vencida")
        self.assertEqual(record["responsavel"], "McLaren")

    def test_response_to_latest_received_cycle_is_assigned_to_transpetro(self):
        record = build_record(
            _document(
                documento="I-PT-4880.00-2493-000-CZ1-001",
                pcf="PCF-I-PT-4880.00-2493-000-CZ1-001_R0D",
                data_pcf=date(2026, 6, 12),
                pcf_resposta="PCF-I-PT-4880.00-2493-000-CZ1-001_R0E",
                data_resposta=date(2026, 6, 15),
            ),
            today=date(2026, 8, 3),
        )

        self.assertTrue(record["resposta_ciclo_atual"])
        self.assertEqual(record["situacao"], "Respondida")
        self.assertEqual(record["responsavel"], "Transpetro")

    def test_excel_export_uses_compatible_formulas(self):
        records = [build_record(_document(), today=date(2026, 2, 10))]
        request = SimpleNamespace(GET={})
        with patch.object(views, "_pcf_response_records", return_value=(records, {})):
            response = views.pcf_controle_respostas_excel.__wrapped__(request)

        with ZipFile(BytesIO(response.content)) as package:
            formulas_xml = b"".join(
                package.read(name)
                for name in package.namelist()
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
            )

        for unsupported in (b"FILTER(", b"TAKE(", b"SORTBY(", b"CHOOSECOLS(", b"AGGREGATE("):
            self.assertNotIn(unsupported, formulas_xml)
        self.assertIn(b"MATCH(", formulas_xml)
        self.assertIn(b"COUNTIF($Y$2:Y", formulas_xml)
        self.assertIn(b"SUM('Dados PCFs'!Y2:Y", formulas_xml)
        self.assertNotIn(b"COUNTA('Base Filtrada'!A2:A", formulas_xml)

    def test_executive_presentation_has_six_slides(self):
        records = [build_record(_document(open_comments=1, qtd_comentarios=1), today=date(2026, 2, 10))]
        payload = build_pcf_executive_presentation(records, summarize(records), date(2026, 2, 10))
        deck = Presentation(BytesIO(payload))

        self.assertEqual(len(deck.slides), 6)
        self.assertIn("LD PROJETO BÁSICO", " ".join(shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text_frame")))
        slide_five_text = " ".join(shape.text for shape in deck.slides[4].shapes if hasattr(shape, "text_frame"))
        slide_six_text = " ".join(shape.text for shape in deck.slides[5].shapes if hasattr(shape, "text_frame"))
        self.assertIn("OPEN / UNDER REVIEW / CLOSED", slide_five_text)
        self.assertNotIn("O/U/C", slide_five_text)
        self.assertIn("TIPOS DE DOCUMENTO", slide_six_text)
        self.assertIn("DISCIPLINAS", slide_six_text)
