from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase
from openpyxl import Workbook

from scripts.build_ld_executive_dashboard import load_dashboard_links


class DashboardDoxLinksTests(SimpleTestCase):
    def test_cadastros_manuais_prevalecem_sobre_exportacoes(self):
        workbook = Workbook()
        fap = workbook.active
        fap.title = "FAP PRODUÇÃO"
        fap.append(["Link", "Documento", "Revisão"])
        fap.append(["https://dox.example/antigo-documento", "I-DE-001", "A"])
        pcfs = workbook.create_sheet("Lista de Documentos_PCFS")
        pcfs.cell(2, 2, "PCF-I-DE-001-R0")
        pcfs.cell(2, 12, "https://dox.example/antiga-pcf")
        with TemporaryDirectory() as temporary:
            source = Path(temporary) / "links.xlsx"
            workbook.save(source)
            dox_links, pcf_links = load_dashboard_links(
                source,
                {("I-DE-001", "A"): "https://dox.novaengevix.com.br/Explorer?d=novo-doc"},
                {"PCF-I-DE-001-R0": "https://dox.novaengevix.com.br/Explorer?d=nova-pcf"},
            )
        self.assertIn("novo-doc", dox_links[("IDE001", "A")])
        self.assertIn("nova-pcf", pcf_links["PCFIDE001R0"])

    def test_reads_exact_fap_revision_and_exact_pcf_links(self):
        with TemporaryDirectory(dir=Path.cwd() / ".test_tmp") as temporary:
            source = Path(temporary) / "links.xlsx"
            workbook = Workbook()
            fap = workbook.active
            fap.title = "FAP Produção (Excel)"
            fap.append(["Documento", "Número Cliente", "Rev"])
            fap.append(["DOX-1", "I-DE-001", "A"])
            fap["A2"].hyperlink = "https://dox.example/documento-a"

            pcfs = workbook.create_sheet("Lista de Documentos_PCFS")
            pcfs.append([None] * 12)
            pcfs.append([None, "Nome", None, "Rótulo", None, None, None, None, None, None, None, "Link"])
            pcfs.cell(3, 2, "PCF-I-DE-001_R0")
            pcfs.cell(3, 12, "https://dox.example/pcf-r0")
            pcfs.cell(4, 2, "PCF-I-DE-001_R0A")
            pcfs.cell(4, 12, "https://dox.example/pcf-r0a")
            workbook.save(source)

            dox_links, pcf_links = load_dashboard_links(source)

        self.assertEqual(dox_links[("IDE001", "A")], "https://dox.example/documento-a")
        self.assertEqual(pcf_links["PCFIDE001R0"], "https://dox.example/pcf-r0")
        self.assertEqual(pcf_links["PCFIDE001R0A"], "https://dox.example/pcf-r0a")

    def test_ignores_non_http_links(self):
        with TemporaryDirectory(dir=Path.cwd() / ".test_tmp") as temporary:
            source = Path(temporary) / "unsafe.xlsx"
            workbook = Workbook()
            fap = workbook.active
            fap.title = "FAP Producao Excel"
            fap.append(["Documento", "Numero Cliente", "Rev"])
            fap.append(["DOX-1", "I-DE-001", 0])
            fap["A2"].hyperlink = "javascript:alert(1)"
            workbook.save(source)

            dox_links, pcf_links = load_dashboard_links(source)

        self.assertEqual(dox_links, {})
        self.assertEqual(pcf_links, {})
