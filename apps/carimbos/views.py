from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from urllib.parse import urlencode
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .forms import CopiaControladaForm
from .models import DistribuicaoCopia
from .services.guia_parser import GuiaParseError
from .services.guia_processor import (
    carregar_previa,
    listar_guias,
    processar_guia,
    raiz_guias,
)
from .services.pdf_stamper import PdfStampError, gerar_copia_controlada
from .services.revision_control import analisar_revisoes
from .services.shopdrawing_distribution import montar_matriz_previa
from apps.contas.permissions import has_perm


def _nome_usuario(request):
    nome = request.user.get_full_name().strip()
    return nome or request.user.get_username()


@has_perm("copias.operar")
def criar_copia_controlada(request):
    if request.method == "POST":
        form = CopiaControladaForm(request.POST, request.FILES)
        if form.is_valid():
            arquivo = form.cleaned_data["arquivo_pdf"]
            try:
                resultado = gerar_copia_controlada(
                    arquivo.read(),
                    nome_original=arquivo.name,
                    numero_gi=form.cleaned_data["numero_gi"],
                    usuario=form.cleaned_data["usuario"],
                    emitido_em=timezone.localtime(),
                )
            except PdfStampError as exc:
                form.add_error("arquivo_pdf", str(exc))
            else:
                response = HttpResponse(
                    resultado.conteudo,
                    content_type="application/pdf",
                )
                response["Content-Disposition"] = (
                    f'attachment; filename="{resultado.nome_arquivo}"'
                )
                response["X-Content-Type-Options"] = "nosniff"
                return response
    else:
        form = CopiaControladaForm(initial={"usuario": _nome_usuario(request)})

    return render(
        request,
        "carimbos/criar_copia_controlada.html",
        {"form": form},
    )


@has_perm("copias.operar")
def guias_emissao(request):
    numero = request.GET.get("guia", "").strip()
    previa = None
    analise_revisoes = None
    destinatarios_exibicao = []
    matriz_shopdrawings = ()
    erro_matriz_shopdrawings = ""
    erro = ""
    if numero:
        try:
            previa = carregar_previa(numero)
            analise_revisoes = analisar_revisoes(previa)
            try:
                matriz_shopdrawings = montar_matriz_previa(previa.documentos)
            except (OSError, KeyError, ValueError) as exc:
                erro_matriz_shopdrawings = (
                    "A matriz recomendada de Shopdrawings está temporariamente "
                    f"indisponível: {exc}"
                )
            destinatarios_exibicao = list(previa.destinatarios)
            existentes = {item.nome.casefold() for item in destinatarios_exibicao}
            from .services.guia_parser import DestinatarioGuia

            for impacto in analise_revisoes.impactos:
                if impacto.destinatario.casefold() not in existentes:
                    destinatarios_exibicao.append(
                        DestinatarioGuia(
                            impacto.destinatario,
                            impacto.email_destinatario,
                            na_guia=False,
                        )
                    )
                    existentes.add(impacto.destinatario.casefold())
        except GuiaParseError as exc:
            erro = str(exc)
    try:
        guias = listar_guias()
    except OSError as exc:
        guias = []
        erro = f"Não foi possível acessar a pasta oficial de GIs/GEs: {exc}"
    return render(
        request,
        "carimbos/guias_emissao.html",
        {
            "raiz_guias": raiz_guias(),
            "guias": guias,
            "numero_selecionado": numero,
            "previa": previa,
            "analise_revisoes": analise_revisoes,
            "destinatarios_exibicao": destinatarios_exibicao,
            "matriz_shopdrawings": matriz_shopdrawings,
            "erro_matriz_shopdrawings": erro_matriz_shopdrawings,
            "erro": erro,
        },
    )


@has_perm("copias.operar")
def processar_guia_emissao(request):
    if request.method != "POST":
        return redirect("carimbos:guias")
    numero = request.POST.get("guia", "").strip()
    destinatarios = request.POST.getlist("destinatarios")
    nomes_carimbo = {}
    meios_distribuicao = {}
    quantidades = {}
    meios = request.POST.getlist("meios_distribuicao")
    quantidades_post = request.POST.getlist("quantidades")
    for indice, nome_original in enumerate(request.POST.getlist("destinatarios_originais")):
        nomes_carimbo[nome_original] = request.POST.getlist("nomes_carimbo")[indice] if indice < len(request.POST.getlist("nomes_carimbo")) else nome_original
        meios_distribuicao[nome_original] = meios[indice] if indice < len(meios) else DistribuicaoCopia.MEIO_NAO_INFORMADO
        quantidades[nome_original] = quantidades_post[indice] if indice < len(quantidades_post) else 1
    try:
        resultado = processar_guia(
            numero,
            destinatarios_selecionados=destinatarios,
            nomes_carimbo=nomes_carimbo,
            meios_distribuicao=meios_distribuicao,
            quantidades=quantidades,
            corrigir_orientacao_paisagem=request.POST.get("corrigir_orientacao_paisagem") == "1",
            posicionar_em_espaco_livre=request.POST.get("posicionar_em_espaco_livre") == "1",
            usar_folha_controle=request.POST.get("usar_folha_controle") == "1",
            usuario=request.user,
        )
    except (GuiaParseError, PdfStampError, OSError) as exc:
        messages.error(request, f"GI/GE não processada: {exc}")
        return redirect(f"{reverse('carimbos:guias')}?{urlencode({'guia': numero})}")

    messages.success(
        request,
        (
            f"GI/GE {resultado.guia.numero} processada: "
            f"{resultado.copias_geradas} novas cópias, "
            f"{resultado.copias_atualizadas} cópias atualizadas e "
            f"{resultado.recolhimentos_abertos} recolhimentos abertos."
        ),
    )
    return redirect(f"{reverse('carimbos:guias')}?{urlencode({'guia': numero})}")


@has_perm("copias.visualizar")
def rastreabilidade(request):
    registros = _filtrar_rastreabilidade(request)
    totais = {
        "registros": registros.count(),
        "quantidade": sum(registros.values_list("quantidade", flat=True)),
        "entregues": registros.filter(status=DistribuicaoCopia.STATUS_ENTREGUE).count(),
        "aguardando": registros.filter(status=DistribuicaoCopia.STATUS_EMITIDA).count(),
    }
    return render(
        request,
        "carimbos/rastreabilidade.html",
        {
            "registros": registros[:500],
            "termo": request.GET.get("q", "").strip(),
            "status_selecionado": request.GET.get("status", "").strip(),
            "meio_selecionado": request.GET.get("meio", "").strip(),
            "status_choices": DistribuicaoCopia.STATUS_CHOICES,
            "meio_choices": DistribuicaoCopia.MEIO_CHOICES,
            "pendentes": DistribuicaoCopia.objects.filter(
                status=DistribuicaoCopia.STATUS_RECOLHIMENTO_PENDENTE
            ).count(),
            "totais": totais,
        },
    )


def _filtrar_rastreabilidade(request):
    termo = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    meio = request.GET.get("meio", "").strip()
    registros = DistribuicaoCopia.objects.select_related(
        "guia", "entregue_por", "recolhida_por"
    )
    if termo:
        registros = registros.filter(
            Q(documento__icontains=termo)
            | Q(guia__numero__icontains=termo)
            | Q(destinatario__icontains=termo)
            | Q(recebedor_carimbo__icontains=termo)
        )
    if status:
        registros = registros.filter(status=status)
    if meio:
        registros = registros.filter(meio_distribuicao=meio)
    return registros


@has_perm("copias.visualizar")
def exportar_rastreabilidade(request):
    registros = _filtrar_rastreabilidade(request)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Controle de Cópias"
    headers = [
        "Documento", "Revisão", "GI/GE", "Data de emissão", "Emitido por",
        "Grupo destinatário", "Pessoa no carimbo", "E-mail", "Meio", "Quantidade",
        "Entrega confirmada em", "Confirmada por", "Situação", "Recolhida em", "Observação",
    ]
    sheet.append(headers)
    for registro in registros:
        sheet.append([
            registro.documento,
            registro.revisao,
            registro.guia.numero,
            timezone.localtime(registro.emitida_em).replace(tzinfo=None),
            registro.guia.remetente,
            registro.destinatario,
            registro.recebedor_carimbo,
            registro.email_destinatario,
            registro.get_meio_distribuicao_display(),
            registro.quantidade,
            timezone.localtime(registro.entregue_em).replace(tzinfo=None) if registro.entregue_em else None,
            registro.entregue_por.get_username() if registro.entregue_por else "",
            registro.get_status_display(),
            timezone.localtime(registro.recolhida_em).replace(tzinfo=None) if registro.recolhida_em else None,
            registro.observacao,
        ])
    header_fill = PatternFill("solid", fgColor="173F54")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [34, 10, 28, 20, 24, 24, 26, 28, 18, 12, 22, 20, 30, 20, 38]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = "dd/mm/yyyy hh:mm"
        sheet.cell(row, 11).number_format = "dd/mm/yyyy hh:mm"
        sheet.cell(row, 14).number_format = "dd/mm/yyyy hh:mm"
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="Controle_Copias_Controladas.xlsx"'
    return response


@has_perm("copias.operar")
def confirmar_entrega(request, pk):
    if request.method != "POST":
        return redirect("carimbos:rastreabilidade")
    registro = get_object_or_404(DistribuicaoCopia, pk=pk)
    registro.status = DistribuicaoCopia.STATUS_ENTREGUE
    registro.entregue_em = timezone.now()
    registro.entregue_por = request.user
    observacao = request.POST.get("observacao", "").strip()
    if observacao:
        registro.observacao = observacao
    registro.save(update_fields=[
        "status", "entregue_em", "entregue_por", "observacao", "atualizado_em",
    ])
    messages.success(request, "Entrega confirmada e registrada na rastreabilidade.")
    return redirect("carimbos:rastreabilidade")


@has_perm("copias.operar")
def confirmar_recolhimento(request, pk):
    if request.method != "POST":
        return redirect("carimbos:rastreabilidade")
    registro = get_object_or_404(DistribuicaoCopia, pk=pk)
    registro.status = DistribuicaoCopia.STATUS_RECOLHIDA
    registro.recolhida_em = timezone.now()
    registro.recolhida_por = request.user
    registro.observacao = request.POST.get("observacao", "").strip()
    registro.save(
        update_fields=[
            "status",
            "recolhida_em",
            "recolhida_por",
            "observacao",
            "atualizado_em",
        ]
    )
    messages.success(request, "Recolhimento confirmado e registrado.")
    return redirect("carimbos:rastreabilidade")
