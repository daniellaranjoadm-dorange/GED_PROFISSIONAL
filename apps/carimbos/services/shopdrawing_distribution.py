from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook

from .guia_parser import DocumentoGuia, normalizar_documento


@dataclass(frozen=True)
class DestinoRecomendado:
    setor: str
    pessoa: str
    meio: str
    quantidade: int = 1
    pendente: bool = False


@dataclass(frozen=True)
class RegraDistribuicao:
    categoria: str
    aliases: tuple[str, ...]
    destinos: tuple[DestinoRecomendado, ...]


@dataclass(frozen=True)
class DocumentoDistribuicao:
    documento: str
    revisao: str
    titulo: str
    categoria: str
    link_dox: str
    destinos: tuple[DestinoRecomendado, ...]
    encontrado_ld: bool
    regra_encontrada: bool


CATEGORIAS = (
    ("Assembly Plan", ("ASSEMBLY",)),
    ("Plate Nesting / Profile Workshop", ("PLATE NESTING", "PROFILE WORKSHOP")),
    ("Part List Plate / Part List Profile", ("PART LIST PLATE", "PART LIST PROFILE")),
    ("Panel Sketch", ("PANEL SKETCH",)),
    (
        "Bill of Material",
        (
            "BILL OF MATERIAL PLATES",
            "BILL OF MATERIAL PROFILES",
            "BILL OF MATERIALS PLATES",
            "BILL OF MATERIALS PROFILES",
        ),
    ),
)


def _normalizar(valor: str) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(char for char in texto if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9]+", " ", texto.upper())).strip()


def _categoria_cabecalho(cabecalho: str):
    normalizado = _normalizar(cabecalho)
    for categoria, aliases in CATEGORIAS:
        if any(alias in normalizado for alias in aliases):
            return categoria, aliases
    return None, ()


def _ler_texto(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def _destinos_linha(linha: str) -> tuple[DestinoRecomendado, ...]:
    if "-" not in linha:
        return ()
    setor, detalhe = (parte.strip() for parte in linha.split("-", 1))
    normalizado = _normalizar(detalhe)
    meio = "DIGITAL" if "DIGITAL" in normalizado else "FISICO"
    pendente = "AINDA NAO TEM" in normalizado
    pessoas = []
    parenteses = re.findall(r"\(([^)]+)\)", detalhe)
    if parenteses and not pendente:
        pessoas.extend(item.strip() for item in parenteses if "cópia" not in item.casefold())
    detalhe_sem_parenteses = re.sub(r"\([^)]*\)", "", detalhe)
    if not pendente:
        for trecho in re.split(r"\s*,\s*", detalhe_sem_parenteses):
            trecho = re.sub(r"\bapenas\s+c[oó]pia.*$", "", trecho, flags=re.IGNORECASE).strip()
            trecho = re.sub(r"\b\d+\s+c[oó]pia.*$", "", trecho, flags=re.IGNORECASE).strip()
            if not trecho or trecho[0].isdigit():
                continue
            pessoas.extend(
                pessoa.strip()
                for pessoa in re.split(r"\s+e\s+", trecho, flags=re.IGNORECASE)
                if pessoa.strip()
            )
    pessoas = list(dict.fromkeys(pessoas))
    if not pessoas:
        return (DestinoRecomendado(setor, "", meio, 1, True),)
    return tuple(DestinoRecomendado(setor, pessoa, meio) for pessoa in pessoas)


@lru_cache(maxsize=8)
def _carregar_regras_cache(path_texto: str, mtime_ns: int) -> tuple[RegraDistribuicao, ...]:
    texto = _ler_texto(Path(path_texto))
    regras = []
    for bloco in re.split(r"\n\s*\n", texto):
        linhas = [linha.strip() for linha in bloco.splitlines() if linha.strip()]
        if len(linhas) < 2:
            continue
        categoria, aliases = _categoria_cabecalho(linhas[0])
        if not categoria:
            continue
        destinos = tuple(destino for linha in linhas[1:] for destino in _destinos_linha(linha))
        regras.append(RegraDistribuicao(categoria, aliases, destinos))
    return tuple(regras)


def carregar_regras(path: Path | None = None) -> tuple[RegraDistribuicao, ...]:
    path = Path(path or settings.SHOPDRAWING_DISTRIBUTION_FILE)
    return _carregar_regras_cache(str(path), path.stat().st_mtime_ns)


@lru_cache(maxsize=8)
def _carregar_indice_cache(path_texto: str, mtime_ns: int):
    workbook = load_workbook(path_texto, read_only=True, data_only=True, keep_links=True)
    sheet = workbook["SD-INDICE"]
    indice = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        documento = str(values[0] or "").strip()
        if not documento:
            continue
        indice[normalizar_documento(documento)] = {
            "titulo": str(values[3] or values[5] or "").strip(),
            "revisao": str(values[7] if values[7] not in (None, "") else "0"),
            "link_dox": str(values[15] or "").strip(),
        }
    workbook.close()
    return indice


def carregar_indice(path: Path | None = None):
    path = Path(path or settings.SHOPDRAWING_LD_FILE)
    return _carregar_indice_cache(str(path), path.stat().st_mtime_ns)


def montar_matriz_previa(documentos: tuple[DocumentoGuia, ...]):
    regras = carregar_regras()
    indice = carregar_indice()
    resultado = []
    for documento in documentos:
        dados = indice.get(normalizar_documento(documento.numero), {})
        titulo = dados.get("titulo", "")
        titulo_normalizado = _normalizar(titulo)
        regra = next(
            (item for item in regras if any(alias in titulo_normalizado for alias in item.aliases)),
            None,
        )
        resultado.append(
            DocumentoDistribuicao(
                documento=documento.numero,
                revisao=documento.revisao,
                titulo=titulo,
                categoria=regra.categoria if regra else "Não classificado",
                link_dox=dados.get("link_dox", ""),
                destinos=regra.destinos if regra else (),
                encontrado_ld=bool(dados),
                regra_encontrada=regra is not None,
            )
        )
    return tuple(resultado)
